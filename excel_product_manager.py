"""Excel price-list uploader and editor for product records."""

import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl


class ExcelProductManagerWindow(tk.Toplevel):
    """Uploads an Excel price list, edits records, and appends products."""

    def __init__(self, parent):
        super().__init__(parent)
        self.title("Excel Price List Manager")
        self.geometry("1100x680")
        self.minsize(850, 550)
        self.transient(parent)
        self.file_path = self.workbook = self.worksheet = None
        self.headers, self.field_vars, self.selected_excel_row = [], {}, None
        self.sheet_var = tk.StringVar()
        self.file_name_var = tk.StringVar(value="No Excel file selected")
        self._build()

    def _build(self):
        top = ttk.Frame(self, padding=12)
        top.pack(fill="x")
        ttk.Button(top, text="Upload Excel File", command=self.choose_file).pack(side="left")
        ttk.Label(top, textvariable=self.file_name_var).pack(side="left", padx=12)
        ttk.Label(top, text="Worksheet:").pack(side="right", padx=(12, 5))
        self.sheet_combo = ttk.Combobox(top, textvariable=self.sheet_var, state="readonly", width=24)
        self.sheet_combo.pack(side="right")
        self.sheet_combo.bind("<<ComboboxSelected>>", lambda _event: self.load_sheet())

        table_box = ttk.LabelFrame(self, text=" Products in Excel Sheet ", padding=8)
        table_box.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        self.tree = ttk.Treeview(table_box, show="headings", selectmode="browse")
        scroll_y = ttk.Scrollbar(table_box, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(table_box, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scroll_y.pack(side="right", fill="y")
        scroll_x.pack(side="bottom", fill="x")
        self.tree.bind("<<TreeviewSelect>>", self.load_selected_product)

        self.editor = ttk.LabelFrame(self, text=" Product Details ", padding=10)
        self.editor.pack(fill="x", padx=12, pady=(0, 12))
        self.editor_inner = ttk.Frame(self.editor)
        self.editor_inner.pack(fill="x")
        ttk.Label(self.editor_inner, text="Upload an Excel file to edit products.").pack(anchor="w")

        actions = ttk.Frame(self, padding=(12, 0, 12, 12))
        actions.pack(fill="x")
        ttk.Button(actions, text="+ New Product", command=self.new_product).pack(side="left")
        ttk.Button(actions, text="Save Product to Excel", command=self.save_product).pack(side="left", padx=8)
        ttk.Button(actions, text="Reload File", command=self.reload_file).pack(side="right")

    def choose_file(self):
        path = filedialog.askopenfilename(parent=self, title="Choose Price List Excel File", filetypes=(("Excel files", "*.xlsx *.xlsm"),))
        if path:
            self.open_file(path)

    def open_file(self, path):
        try:
            self.file_path = path
            self.workbook = openpyxl.load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
            self.file_name_var.set(os.path.basename(path))
            self.sheet_combo["values"] = self.workbook.sheetnames
            self.sheet_var.set(self.workbook.active.title)
            self.load_sheet()
        except Exception as error:
            messagebox.showerror("Cannot Open Excel", f"Could not open this Excel file:\n{error}", parent=self)

    def load_sheet(self):
        if not self.workbook:
            return
        self.worksheet = self.workbook[self.sheet_var.get()]
        self.headers = [str(cell.value).strip() if cell.value is not None else "" for cell in self.worksheet[1]]
        used_headers = [header for header in self.headers if header]
        if not used_headers:
            messagebox.showwarning("Missing Headers", "The first row of this worksheet must contain column names.", parent=self)
            return
        if len({header.lower() for header in used_headers}) != len(used_headers):
            messagebox.showwarning("Duplicate Headers", "Each first-row column name must be unique.", parent=self)
            return
        self.build_editor()
        self.refresh_table()
        self.new_product()

    def build_editor(self):
        for widget in self.editor_inner.winfo_children():
            widget.destroy()
        self.field_vars = {header: tk.StringVar() for header in self.headers if header}
        for index, header in enumerate(self.field_vars):
            row, group = divmod(index, 2)
            column = group * 2
            ttk.Label(self.editor_inner, text=f"{header}:").grid(row=row, column=column, sticky="w", padx=(0, 6), pady=4)
            ttk.Entry(self.editor_inner, textvariable=self.field_vars[header], width=34).grid(row=row, column=column + 1, sticky="ew", padx=(0, 18), pady=4)
        self.editor_inner.columnconfigure(1, weight=1)
        self.editor_inner.columnconfigure(3, weight=1)

    def refresh_table(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        columns = list(self.field_vars)
        self.tree["columns"] = columns
        for column in columns:
            self.tree.heading(column, text=column)
            self.tree.column(column, width=150, minwidth=90, anchor="w")
        for excel_row in range(2, self.worksheet.max_row + 1):
            values = [self.worksheet.cell(excel_row, self.headers.index(column) + 1).value for column in columns]
            if any(value not in (None, "") for value in values):
                self.tree.insert("", "end", iid=str(excel_row), values=["" if value is None else str(value) for value in values])

    def load_selected_product(self, _event=None):
        selected = self.tree.selection()
        if not selected:
            return
        self.selected_excel_row = int(selected[0])
        for column, variable in self.field_vars.items():
            value = self.worksheet.cell(self.selected_excel_row, self.headers.index(column) + 1).value
            variable.set("" if value is None else str(value))

    def new_product(self):
        self.selected_excel_row = None
        for variable in self.field_vars.values():
            variable.set("")
        for item in self.tree.selection():
            self.tree.selection_remove(item)

    def save_product(self):
        if not self.worksheet:
            messagebox.showwarning("Upload Needed", "Upload an Excel price list first.", parent=self)
            return
        if not any(variable.get().strip() for variable in self.field_vars.values()):
            messagebox.showwarning("Product Details", "Enter at least one product detail before saving.", parent=self)
            return
        excel_row = self.selected_excel_row or self.worksheet.max_row + 1
        for column, variable in self.field_vars.items():
            self.worksheet.cell(excel_row, self.headers.index(column) + 1).value = variable.get().strip()
        try:
            self.workbook.save(self.file_path)
            self.refresh_table()
            self.selected_excel_row = excel_row
            messagebox.showinfo("Saved", "Product details were saved to the Excel file.", parent=self)
        except PermissionError:
            messagebox.showerror("File Is Open", "Close the Excel file in Microsoft Excel, then save again.", parent=self)
        except Exception as error:
            messagebox.showerror("Save Failed", f"Could not save the Excel file:\n{error}", parent=self)

    def reload_file(self):
        if self.file_path:
            current_sheet = self.sheet_var.get()
            self.open_file(self.file_path)
            if current_sheet in self.workbook.sheetnames:
                self.sheet_var.set(current_sheet)
                self.load_sheet()
