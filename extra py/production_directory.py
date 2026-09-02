import csv
import datetime
import glob
import os
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import auth_manager

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from PIL import Image, ImageTk
    HAS_PIL = True
except ImportError:
    HAS_PIL = False


class ProductTestingWindow(tk.Toplevel):
    """Sub-window for Product Testing and Visual Inspection (With Edit & User Tracking)."""
    def __init__(self, parent, batch_folder_path=None):
        super().__init__(parent)
        self.title("Product Testing & Visual Inspection")
        self.geometry("550x650")
        self.transient(parent)
        self.grab_set()

        self.batch_folder_path = batch_folder_path
        self.product_id_var = tk.StringVar()
        self.front_photo_path = None
        self.back_photo_path = None

        self.checks = {
            "No Scratches / Surface Defects": tk.BooleanVar(value=False),
            "Correct Labeling & Barcode": tk.BooleanVar(value=False),
            "Dimensions & Fit Verified": tk.BooleanVar(value=False),
            "Secure Components / Assembly": tk.BooleanVar(value=False)
        }

        self.create_widgets()

    def create_widgets(self):
        main_frame = ttk.Frame(self, padding=15)
        main_frame.pack(fill="both", expand=True)

        id_frame = ttk.LabelFrame(main_frame, text=" Product Identification & Search ", padding=12)
        id_frame.pack(fill="x", pady=(0, 10))

        ttk.Label(id_frame, text="Enter Product ID:", font=("Helvetica", 10, "bold")).grid(row=0, column=0, sticky="w", padx=5, pady=5)
        ttk.Entry(id_frame, textvariable=self.product_id_var, width=20).grid(row=0, column=1, sticky="w", padx=5, pady=5)
        ttk.Button(id_frame, text="🔍 Search / Edit", command=self.search_product_record).grid(row=0, column=2, padx=10, pady=5)

        chk_frame = ttk.LabelFrame(main_frame, text=" Visual Inspection Checklist ", padding=12)
        chk_frame.pack(fill="x", pady=(0, 10))

        for text, var in self.checks.items():
            ttk.Checkbutton(chk_frame, text=text, variable=var).pack(anchor="w", pady=4, padx=5)

        photo_frame = ttk.LabelFrame(main_frame, text=" Individual Product Photos ", padding=12)
        photo_frame.pack(fill="x", pady=(0, 15))

        ttk.Button(photo_frame, text="Upload Front Photo", command=lambda: self.load_photo("front")).grid(row=0, column=0, padx=5, pady=5)
        self.lbl_front = ttk.Label(photo_frame, text="No front photo selected", font=("Helvetica", 9, "italic"))
        self.lbl_front.grid(row=0, column=1, sticky="w", padx=5, pady=5)

        ttk.Button(photo_frame, text="Upload Back Photo", command=lambda: self.load_photo("back")).grid(row=1, column=0, padx=5, pady=5)
        self.lbl_back = ttk.Label(photo_frame, text="No back photo selected", font=("Helvetica", 9, "italic"))
        self.lbl_back.grid(row=1, column=1, sticky="w", padx=5, pady=5)

        self.btn_save = ttk.Button(main_frame, text="💾 Save / Update Inspection Record", command=self.save_inspection_data)
        self.btn_save.pack(pady=10)

    def load_photo(self, photo_type):
        file_path = filedialog.askopenfilename(
            title=f"Select {photo_type.capitalize()} View Photo",
            filetypes=[("Image Files", "*.jpg *.jpeg *.png *.bmp")],
            parent=self
        )
        if file_path:
            if photo_type == "front":
                self.front_photo_path = file_path
                self.lbl_front.config(text=os.path.basename(file_path))
            else:
                self.back_photo_path = file_path
                self.lbl_back.config(text=os.path.basename(file_path))

    def search_product_record(self):
        prod_id = self.product_id_var.get().strip()
        if not prod_id:
            messagebox.showwarning("Validation Error", "Please enter a Product ID to search.", parent=self)
            return

        base_target_dir = self.batch_folder_path if self.batch_folder_path and os.path.exists(self.batch_folder_path) else os.path.join(os.getcwd(), "Production")
        csv_file = os.path.join(base_target_dir, "product_testing_log.csv")

        record_found = False
        if os.path.exists(csv_file):
            try:
                with open(csv_file, mode="r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    for row in reader:
                        if row.get("Product_ID") == prod_id:
                            record_found = True
                            for text, var in self.checks.items():
                                if text in row:
                                    var.set(row[text] == "Passed")
                            self.lbl_front.config(text=row.get("Front_Photo", "No front photo selected"))
                            self.lbl_back.config(text=row.get("Back_Photo", "No back photo selected"))
                            break
            except Exception as e:
                messagebox.showerror("Read Error", f"Failed to search record:\n{e}", parent=self)
                return

        if record_found:
            messagebox.showinfo("Record Found", f"Existing record for '{prod_id}' loaded.", parent=self)
        else:
            messagebox.showinfo("Not Found", f"No existing record found for '{prod_id}'.", parent=self)

    def get_current_user_name(self):
        current_user = getattr(auth_manager, "current_user", {})
        user_name = "Admin"
        if isinstance(current_user, dict):
            user_name = current_user.get("full_name", current_user.get("username", "Admin"))
        elif hasattr(current_user, "username"):
            user_name = getattr(current_user, "full_name", getattr(current_user, "username", "Admin"))
        return user_name

    def save_inspection_data(self):
        prod_id = self.product_id_var.get().strip()
        if not prod_id:
            messagebox.showwarning("Validation Error", "Please enter a valid Product ID.", parent=self)
            return

        base_target_dir = self.batch_folder_path if self.batch_folder_path and os.path.exists(self.batch_folder_path) else os.path.join(os.getcwd(), "Production")
        prod_folder = os.path.join(base_target_dir, "Product_Testing_Records", prod_id)
        os.makedirs(prod_folder, exist_ok=True)

        saved_front = "None"
        saved_back = "None"

        try:
            if self.front_photo_path and os.path.exists(self.front_photo_path):
                ext = os.path.splitext(self.front_photo_path)[1]
                saved_front = f"front_view{ext}"
                dest_front = os.path.join(prod_folder, saved_front)
                if HAS_PIL:
                    Image.open(self.front_photo_path).save(dest_front)
                else:
                    shutil.copy(self.front_photo_path, dest_front)

            if self.back_photo_path and os.path.exists(self.back_photo_path):
                ext = os.path.splitext(self.back_photo_path)[1]
                saved_back = f"back_view{ext}"
                dest_back = os.path.join(prod_folder, saved_back)
                if HAS_PIL:
                    Image.open(self.back_photo_path).save(dest_back)
                else:
                    shutil.copy(self.back_photo_path, dest_back)

            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            action_user = self.get_current_user_name()

            record = {
                "Timestamp": now_str,
                "Product_ID": prod_id,
            }
            for text, var in self.checks.items():
                record[text] = "Passed" if var.get() else "Failed"
            
            record["Front_Photo"] = saved_front if saved_front != "None" else self.lbl_front.cget("text")
            record["Back_Photo"] = saved_back if saved_back != "None" else self.lbl_back.cget("text")
            record["Last_Edited_By"] = action_user

            csv_file = os.path.join(base_target_dir, "product_testing_log.csv")
            rows = []
            updated = False
            fieldnames = list(record.keys())
            
            if os.path.exists(csv_file):
                with open(csv_file, mode="r", newline="", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    for fld in (reader.fieldnames or []):
                        if fld not in fieldnames:
                            fieldnames.append(fld)
                    for row in reader:
                        if row.get("Product_ID") == prod_id:
                            row.update(record)
                            updated = True
                        rows.append(row)

            if not updated:
                rows.append(record)

            with open(csv_file, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)

            if HAS_OPENPYXL:
                excel_file = os.path.join(base_target_dir, "Product_Testing_Report.xlsx")
                if os.path.exists(excel_file):
                    wb = openpyxl.load_workbook(excel_file)
                    ws = wb.active
                else:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Testing Logs"
                    ws.append(fieldnames)

                excel_headers = [cell.value for cell in ws[1]]
                for fld in fieldnames:
                    if fld not in excel_headers:
                        excel_headers.append(fld)
                        ws.cell(row=1, column=len(excel_headers), value=fld)

                prod_id_col_idx = excel_headers.index("Product_ID") + 1 if "Product_ID" in excel_headers else 2
                row_idx_to_update = None
                for row_idx in range(2, ws.max_row + 1):
                    if str(ws.cell(row=row_idx, column=prod_id_col_idx).value) == prod_id:
                        row_idx_to_update = row_idx
                        break

                if row_idx_to_update:
                    for key, val in record.items():
                        col_idx = excel_headers.index(key) + 1
                        ws.cell(row=row_idx_to_update, column=col_idx, value=val)
                else:
                    new_row = [record.get(h, "") for h in excel_headers]
                    ws.append(new_row)

                wb.save(excel_file)

            messagebox.showinfo("Success", f"Testing record saved successfully!\nSaved By: {action_user}", parent=self)
            self.destroy()

        except Exception as e:
            messagebox.showerror("Save Error", f"Failed to save testing data:\n{e}", parent=self)


class ExistingProductionView(ttk.Frame):
    """File 1: Manages Sections 1-4 (Batch Selection, Material List/PO, Deliver & Receive PCB) plus Testing."""
    def __init__(self, parent, back_callback):
        super().__init__(parent)
        self.back_callback = back_callback

        self.prod_dir = os.path.join(os.getcwd(), "Production")
        self.new_prod_dir = os.path.join(self.prod_dir, "New Production Start")

        self.current_materials = []
        self.batch_map = {}
        
        self.is_mat_frame_expanded = False
        self.is_pcb_frame_expanded = False
        self.is_receive_frame_expanded = False

        self.show_existing_production_ui()

    def resequence_sr_no(self):
        for idx, item in enumerate(self.current_materials, start=1):
            item["sr_no"] = str(idx)

    def show_existing_production_ui(self):
        for child in self.winfo_children():
            child.destroy()

        top_bar = ttk.Frame(self, padding=10)
        top_bar.pack(fill="x")

        ttk.Button(top_bar, text="⬅ Back to Option", command=self.back_callback).pack(side="left")
        ttk.Label(top_bar, text="Existing Production Directory", font=("Helvetica", 13, "bold")).pack(side="left", padx=15)

        main_canvas = tk.Canvas(self, borderwidth=0, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=main_canvas.yview)
        scrollable_frame = ttk.Frame(main_canvas, padding=15)

        scrollable_frame.bind("<Configure>", lambda e: main_canvas.configure(scrollregion=main_canvas.bbox("all")))
        main_canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        main_canvas.configure(yscrollcommand=scrollbar.set)

        main_canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # 1. Select Production Batch
        drop_frame = ttk.LabelFrame(scrollable_frame, text=" 1. Select Production Batch ", padding=15)
        drop_frame.pack(fill="x", pady=(0, 15))

        ttk.Label(drop_frame, text="Select Batch:", font=("Helvetica", 10, "bold")).grid(row=0, column=0, sticky="w", padx=5, pady=5)
        self.combo_batch = ttk.Combobox(drop_frame, width=45, state="readonly")
        self.combo_batch.grid(row=0, column=1, sticky="w", padx=10, pady=5)
        self.combo_batch.bind("<<ComboboxSelected>>", self.on_batch_dropdown_selected)

        self.lbl_batch_info = ttk.Label(drop_frame, text="Select a batch from the dropdown above.", font=("Helvetica", 9, "italic"))
        self.lbl_batch_info.grid(row=1, column=0, columnspan=3, sticky="w", padx=5, pady=5)

        # 2. Material List & Purchase Orders
        self.mat_frame = ttk.LabelFrame(scrollable_frame, text=" 2. Material List & Purchase Orders ", padding=10)
        self.mat_frame.pack(fill="both", expand=True, pady=(0, 15))

        mat_header_bar = ttk.Frame(self.mat_frame)
        mat_header_bar.pack(fill="x", side="top", pady=(0, 5))
        self.btn_toggle_mat = ttk.Button(mat_header_bar, text="▼ Expand Section", width=18, command=self.toggle_material_section)
        self.btn_toggle_mat.pack(side="right")

        self.mat_content_container = ttk.Frame(self.mat_frame)
        self.mat_content_container.pack(fill="both", expand=True)
        if not self.is_mat_frame_expanded:
            self.mat_content_container.pack_forget()

        mat_cols = ("PO_Select", "Sr_No", "Category", "Part_Name", "Value", "Package", "Tolerance", "Quantity", "QTY", "Rate", "Total_Price", "Make")
        self.mat_tree = ttk.Treeview(self.mat_content_container, columns=mat_cols, show="headings", selectmode="browse", height=8)
        for col in mat_cols:
            self.mat_tree.heading(col, text=col)
        self.mat_tree.pack(side="top", fill="both", expand=True)
        self.mat_tree.bind("<Button-1>", self.on_material_tree_click)

        mat_btn_bar = ttk.Frame(self.mat_content_container, padding=(0, 10, 0, 0))
        mat_btn_bar.pack(fill="x", side="bottom")
        ttk.Button(mat_btn_bar, text="☑ Select All PO", command=self.toggle_select_all_po).pack(side="left", padx=4)
        ttk.Button(mat_btn_bar, text="📊 Import Excel", command=self.import_and_convert_excel).pack(side="left", padx=4)
        ttk.Button(mat_btn_bar, text="➕ Add Row", command=self.open_add_material_dialog).pack(side="left", padx=4)
        ttk.Button(mat_btn_bar, text="❌ Delete", command=self.remove_selected_material).pack(side="left", padx=4)
        ttk.Button(mat_btn_bar, text="📦 Generate PO", command=self.send_purchase_order).pack(side="left", padx=4)
        ttk.Button(mat_btn_bar, text="💾 Save CSV", command=self.save_material_list_changes).pack(side="left", padx=4)

        # 3. Deliver to PCB Mounting
        self.pcb_frame = ttk.LabelFrame(scrollable_frame, text=" 3. Deliver to PCB Mounting ", padding=10)
        self.pcb_frame.pack(fill="both", expand=True, pady=(0, 15))

        pcb_header_bar = ttk.Frame(self.pcb_frame)
        pcb_header_bar.pack(fill="x", side="top", pady=(0, 5))
        self.btn_toggle_pcb = ttk.Button(pcb_header_bar, text="▼ Expand Section", width=18, command=self.toggle_pcb_section)
        self.btn_toggle_pcb.pack(side="right")

        self.pcb_content_container = ttk.Frame(self.pcb_frame)
        self.pcb_content_container.pack(fill="both", expand=True)
        if not self.is_pcb_frame_expanded:
            self.pcb_content_container.pack_forget()

        pcb_cols = ("PCB_Select", "Sr_No", "Category", "Part_Name", "Value", "Package", "Tolerance", "Quantity", "QTY", "Make")
        self.pcb_tree = ttk.Treeview(self.pcb_content_container, columns=pcb_cols, show="headings", selectmode="browse", height=8)
        for col in pcb_cols:
            self.pcb_tree.heading(col, text=col)
        self.pcb_tree.pack(side="top", fill="both", expand=True)
        self.pcb_tree.bind("<Button-1>", self.on_pcb_tree_click)

        pcb_btn_bar = ttk.Frame(self.pcb_content_container, padding=(0, 10, 0, 0))
        pcb_btn_bar.pack(fill="x", side="bottom")
        ttk.Button(pcb_btn_bar, text="☑ Select All", command=self.toggle_select_all_pcb).pack(side="left", padx=4)
        ttk.Button(pcb_btn_bar, text="📦 Generate List", command=self.send_pcb_material_list).pack(side="left", padx=4)

        # 4. Receive from PCB Mounting
        self.receive_frame = ttk.LabelFrame(scrollable_frame, text=" 4. Receive from PCB Mounting ", padding=10)
        self.receive_frame.pack(fill="both", expand=True, pady=(0, 15))

        receive_header_bar = ttk.Frame(self.receive_frame)
        receive_header_bar.pack(fill="x", side="top", pady=(0, 5))
        self.btn_toggle_receive = ttk.Button(receive_header_bar, text="▼ Expand Section", width=18, command=self.toggle_receive_section)
        self.btn_toggle_receive.pack(side="right")

        self.receive_content_container = ttk.Frame(self.receive_frame)
        self.receive_content_container.pack(fill="both", expand=True)
        if not self.is_receive_frame_expanded:
            self.receive_content_container.pack_forget()

        receive_cols = ("Receive_Select", "Sr_No", "Category", "Part_Name", "Value", "Package", "Tolerance", "Quantity", "QTY", "Make")
        self.receive_tree = ttk.Treeview(self.receive_content_container, columns=receive_cols, show="headings", selectmode="browse", height=8)
        for col in receive_cols:
            self.receive_tree.heading(col, text=col)
        self.receive_tree.pack(side="top", fill="both", expand=True)
        self.receive_tree.bind("<Button-1>", self.on_receive_tree_click)

        receive_btn_bar = ttk.Frame(self.receive_content_container, padding=(0, 10, 0, 0))
        receive_btn_bar.pack(fill="x", side="bottom")
        ttk.Button(receive_btn_bar, text="☑ Select All", command=self.toggle_select_all_receive).pack(side="left", padx=4)
        ttk.Button(receive_btn_bar, text="📦 Generate Received List", command=self.send_receive_material_list).pack(side="left", padx=4)

        # 5. Product Testing Option Window Link
        test_frame = ttk.LabelFrame(scrollable_frame, text=" 5. Product Testing & Visual Inspection ", padding=12)
        test_frame.pack(fill="x", expand=True)
        ttk.Button(test_frame, text="🔍 Open Product Testing Window", command=self.open_product_testing_window).pack(side="left", padx=4)

        self.load_batches_into_dropdown()

    def toggle_material_section(self):
        if self.is_mat_frame_expanded:
            self.mat_content_container.pack_forget()
            self.btn_toggle_mat.config(text="▼ Expand Section")
            self.is_mat_frame_expanded = False
        else:
            self.mat_content_container.pack(fill="both", expand=True)
            self.btn_toggle_mat.config(text="▲ Fold Section")
            self.is_mat_frame_expanded = True

    def toggle_pcb_section(self):
        if self.is_pcb_frame_expanded:
            self.pcb_content_container.pack_forget()
            self.btn_toggle_pcb.config(text="▼ Expand Section")
            self.is_pcb_frame_expanded = False
        else:
            self.pcb_content_container.pack(fill="both", expand=True)
            self.btn_toggle_pcb.config(text="▲ Fold Section")
            self.is_pcb_frame_expanded = True

    def toggle_receive_section(self):
        if self.is_receive_frame_expanded:
            self.receive_content_container.pack_forget()
            self.btn_toggle_receive.config(text="▼ Expand Section")
            self.is_receive_frame_expanded = False
        else:
            self.receive_content_container.pack(fill="both", expand=True)
            self.btn_toggle_receive.config(text="▲ Fold Section")
            self.is_receive_frame_expanded = True

    def open_product_testing_window(self):
        if not hasattr(self, "active_csv_path") or not self.active_csv_path:
            messagebox.showwarning("Batch Needed", "Please select a production batch first.", parent=self)
            return
        batch_folder = os.path.dirname(self.active_csv_path)
        ProductTestingWindow(self, batch_folder_path=batch_folder)

    def load_batches_into_dropdown(self):
        self.batch_map = {}
        batch_options = []
        csv_files = glob.glob(os.path.join(self.new_prod_dir, "*", "info.csv"))

        for csv_file in csv_files:
            try:
                with open(csv_file, mode="r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    row = next(reader, None)
                    if row:
                        p_name = row.get("Product_Name", "Unknown")
                        b_num = row.get("Batch_Number", "Unknown")
                        display_str = f"{b_num} ({p_name})"
                        self.batch_map[display_str] = (csv_file, row)
                        batch_options.append(display_str)
            except Exception as e:
                print(f"Error reading {csv_file}: {e}")

        self.combo_batch["values"] = batch_options
        if batch_options:
            self.combo_batch.set(batch_options[0])
            self.on_batch_dropdown_selected(None)

    def on_batch_dropdown_selected(self, event):
        selected_key = self.combo_batch.get()
        if not selected_key or selected_key not in self.batch_map:
            return
        csv_path, row_data = self.batch_map[selected_key]
        self.active_csv_path = csv_path
        self.active_batch_row = row_data

        exp_date = row_data.get("Expected_Completion_Date", "N/A")
        start_date = row_data.get("Start_Date", "N/A")
        personnel = row_data.get("Assigned_Personnel", "None")

        self.lbl_batch_info.config(text=f"Start Date: {start_date}   |   Expected Completion: {exp_date}   |   Personnel: {personnel}")
        self.refresh_material_tree_from_data(row_data.get("Material_List", ""))

    def refresh_material_tree_from_data(self, raw_materials):
        self.mat_tree.delete(*self.mat_tree.get_children())
        self.pcb_tree.delete(*self.pcb_tree.get_children())
        self.receive_tree.delete(*self.receive_tree.get_children())
        self.current_materials = []

        if raw_materials:
            items = raw_materials.split(";")
            for idx, item_str in enumerate(items, 1):
                parts = item_str.split("|")
                if len(parts) >= 11:
                    po_val = parts[0] == "1"
                    m_dict = {
                        "po_check": po_val, "pcb_check": po_val, "receive_check": po_val,
                        "sr_no": parts[1] or str(idx), "category": parts[2], "part_name": parts[3],
                        "value": parts[4], "package": parts[5], "tolerance": parts[6],
                        "quantity": parts[7], "qty": parts[8], "rate": parts[9], "make": parts[10]
                    }
                else:
                    m_dict = {
                        "po_check": False, "pcb_check": False, "receive_check": False,
                        "sr_no": str(idx), "category": parts[0] if len(parts) > 0 else "",
                        "part_name": parts[1] if len(parts) > 1 else "", "value": parts[2] if len(parts) > 2 else "",
                        "package": parts[3] if len(parts) > 3 else "", "tolerance": "1%",
                        "quantity": parts[4] if len(parts) > 4 else "1", "qty": "", "rate": "", "make": ""
                    }
                self.current_materials.append(m_dict)

        self.resequence_sr_no()
        self.redraw_material_tree()
        self.redraw_pcb_tree()
        self.redraw_receive_tree()

    def redraw_material_tree(self):
        self.mat_tree.delete(*self.mat_tree.get_children())
        for idx, m in enumerate(self.current_materials):
            chk = "☑" if m["po_check"] else "☐"
            try:
                tot = f"{float(m['qty']) * float(m['rate']):.2f}" if m["qty"] and m["rate"] else ""
            except ValueError:
                tot = ""
            self.mat_tree.insert("", "end", values=(chk, m["sr_no"], m["category"], m["part_name"], m["value"], m["package"], m["tolerance"], m["quantity"], m["qty"], m["rate"], tot, m["make"]))

    def redraw_pcb_tree(self):
        self.pcb_tree.delete(*self.pcb_tree.get_children())
        for idx, m in enumerate(self.current_materials):
            chk = "☑" if m.get("pcb_check", False) else "☐"
            self.pcb_tree.insert("", "end", values=(chk, m["sr_no"], m["category"], m["part_name"], m["value"], m["package"], m["tolerance"], m["quantity"], m["qty"], m["make"]))

    def redraw_receive_tree(self):
        self.receive_tree.delete(*self.receive_tree.get_children())
        for idx, m in enumerate(self.current_materials):
            chk = "☑" if m.get("receive_check", False) else "☐"
            self.receive_tree.insert("", "end", values=(chk, m["sr_no"], m["category"], m["part_name"], m["value"], m["package"], m["tolerance"], m["quantity"], m["qty"], m["make"]))

    def toggle_select_all_po(self):
        if not self.current_materials: return
        state = not all(m["po_check"] for m in self.current_materials)
        for m in self.current_materials: m["po_check"] = state
        self.redraw_material_tree()

    def toggle_select_all_pcb(self):
        if not self.current_materials: return
        state = not all(m.get("pcb_check", False) for m in self.current_materials)
        for m in self.current_materials: m["pcb_check"] = state
        self.redraw_pcb_tree()

    def toggle_select_all_receive(self):
        if not self.current_materials: return
        state = not all(m.get("receive_check", False) for m in self.current_materials)
        for m in self.current_materials: m["receive_check"] = state
        self.redraw_receive_tree()

    def on_material_tree_click(self, event):
        if self.mat_tree.identify_region(event.x, event.y) == "cell" and self.mat_tree.identify_column(event.x) == "#1":
            item = self.mat_tree.identify_row(event.y)
            if item:
                sr = str(self.mat_tree.item(item, "values")[1])
                for m in self.current_materials:
                    if str(m["sr_no"]) == sr: m["po_check"] = not m["po_check"]; break
                self.redraw_material_tree()

    def on_pcb_tree_click(self, event):
        if self.pcb_tree.identify_region(event.x, event.y) == "cell" and self.pcb_tree.identify_column(event.x) == "#1":
            item = self.pcb_tree.identify_row(event.y)
            if item:
                sr = str(self.pcb_tree.item(item, "values")[1])
                for m in self.current_materials:
                    if str(m["sr_no"]) == sr: m["pcb_check"] = not m.get("pcb_check", False); break
                self.redraw_pcb_tree()

    def on_receive_tree_click(self, event):
        if self.receive_tree.identify_region(event.x, event.y) == "cell" and self.receive_tree.identify_column(event.x) == "#1":
            item = self.receive_tree.identify_row(event.y)
            if item:
                sr = str(self.receive_tree.item(item, "values")[1])
                for m in self.current_materials:
                    if str(m["sr_no"]) == sr: m["receive_check"] = not m.get("receive_check", False); break
                self.redraw_receive_tree()

    def import_and_convert_excel(self):
        if not HAS_OPENPYXL: return
        file_path = filedialog.askopenfilename(title="Select Excel", filetypes=[("Excel Files", "*.xlsx")], parent=self)
        if not file_path: return
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            current_category = ""
            for sheet_name in wb.sheetnames:
                for row in wb[sheet_name].iter_rows(values_only=True):
                    if not row or all(v is None for v in row): continue
                    r_sr = str(row[0]).strip() if row[0] is not None else ""
                    if r_sr.lower() in ("sr. no.", "sr no"): continue
                    if row[1]: current_category = str(row[1]).strip()
                    self.current_materials.append({
                        "po_check": False, "pcb_check": False, "receive_check": False, "sr_no": "",
                        "category": current_category, "part_name": str(row[2]) if len(row) > 2 and row[2] else "",
                        "value": str(row[3]) if len(row) > 3 and row[3] else "", "package": str(row[4]) if len(row) > 4 and row[4] else "",
                        "tolerance": str(row[5]) if len(row) > 5 and row[5] else "1%", "quantity": str(row[6]) if len(row) > 6 and row[6] else "",
                        "qty": str(row[7]) if len(row) > 7 and row[7] else "", "rate": str(row[8]) if len(row) > 8 and row[8] else "",
                        "make": str(row[9]) if len(row) > 9 and row[9] else ""
                    })
            self.resequence_sr_no()
            self.redraw_material_tree(); self.redraw_pcb_tree(); self.redraw_receive_tree()
            messagebox.showinfo("Success", "Imported successfully!", parent=self)
        except Exception as e:
            messagebox.showerror("Error", f"{e}", parent=self)

    def open_add_material_dialog(self):
        dialog = tk.Toplevel(self)
        dialog.title("Add Material")
        dialog.geometry("380x450")
        fields = ["Category", "Part Name", "Value", "Package", "Tolerance", "Quantity", "QTY", "Rate", "Make"]
        entries = {}
        for idx, f in enumerate(fields):
            ttk.Label(dialog, text=f"{f}:").grid(row=idx, column=0, sticky="w", padx=15, pady=4)
            ent = ttk.Entry(dialog, width=25)
            ent.grid(row=idx, column=1, padx=15, pady=4)
            entries[f] = ent
        def save_item():
            self.current_materials.append({
                "po_check": False, "pcb_check": False, "receive_check": False,
                "sr_no": str(len(self.current_materials) + 1), "category": entries["Category"].get().strip(),
                "part_name": entries["Part Name"].get().strip(), "value": entries["Value"].get().strip(),
                "package": entries["Package"].get().strip(), "tolerance": entries["Tolerance"].get().strip(),
                "quantity": entries["Quantity"].get().strip(), "qty": entries["QTY"].get().strip(),
                "rate": entries["Rate"].get().strip(), "make": entries["Make"].get().strip()
            })
            self.resequence_sr_no(); self.redraw_material_tree(); self.redraw_pcb_tree(); self.redraw_receive_tree()
            dialog.destroy()
        ttk.Button(dialog, text="Add", command=save_item).grid(row=len(fields), column=0, columnspan=2, pady=15)

    def remove_selected_material(self):
        for tree, name in [(self.mat_tree, "mat"), (self.pcb_tree, "pcb"), (self.receive_tree, "rcv")]:
            sel = tree.selection()
            if sel:
                del self.current_materials[tree.index(sel[0])]
                self.resequence_sr_no()
                self.redraw_material_tree(); self.redraw_pcb_tree(); self.redraw_receive_tree()
                return
        messagebox.showwarning("Warning", "Select a row to delete.", parent=self)

    def save_material_list_changes(self):
        if not hasattr(self, "active_csv_path"): return
        self.resequence_sr_no()
        mat_str_list = [f"{'1' if m['po_check'] else '0'}|{m['sr_no']}|{m['category']}|{m['part_name']}|{m['value']}|{m['package']}|{m['tolerance']}|{m['quantity']}|{m['qty']}|{m['rate']}|{m['make']}" for m in self.current_materials]
        self.active_batch_row["Material_List"] = ";".join(mat_str_list)
        try:
            with open(self.active_csv_path, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=list(self.active_batch_row.keys()))
                writer.writeheader(); writer.writerow(self.active_batch_row)
            messagebox.showinfo("Success", "Saved successfully!", parent=self)
        except Exception as e:
            messagebox.showerror("Error", f"{e}", parent=self)

    def send_purchase_order(self):
        self._export_excel_subset("po_check", "Purchase_Order")

    def send_pcb_material_list(self):
        self._export_excel_subset("pcb_check", "PCB_Mounting_Material_List")

    def send_receive_material_list(self):
        self._export_excel_subset("receive_check", "Received_from_PCB_Mounting")

    def _export_excel_subset(self, check_key, prefix):
        if not HAS_OPENPYXL or not hasattr(self, "active_csv_path"): return
        items = [m for m in self.current_materials if m.get(check_key, False)]
        if not items:
            messagebox.showwarning("Warning", "No items selected.", parent=self)
            return
        try:
            path = os.path.join(os.path.dirname(self.active_csv_path), f"{prefix}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append([prefix.replace("_", " ").upper()])
            headers = ["Sr. No.", "Category", "Part Name", "Value", "Package", "Tolerance", "Quantity", "PO QTY", "Make"]
            ws.append([]); ws.append(headers)
            start = ws.max_row + 1
            for idx, m in enumerate(items, 1):
                ws.append([str(idx), m.get("category", ""), m.get("part_name", ""), m.get("value", ""), m.get("package", ""), m.get("tolerance", ""), m.get("quantity", ""), m.get("qty", ""), m.get("make", "")])
            wb.save(path)
            messagebox.showinfo("Success", f"Generated successfully:\n{path}", parent=self)
        except Exception as e:
            messagebox.showerror("Error", f"{e}", parent=self)