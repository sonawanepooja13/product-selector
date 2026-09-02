import csv
import datetime
import glob
import os
import shutil
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import auth_manager

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from PIL import Image, ImageTk
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import cv2
    HAS_CV2 = True
except ImportError:
    HAS_CV2 = False


class ProductTestingWindow(tk.Toplevel):
    """Sub-window for Product Testing and Visual Inspection (Fully Responsive Layout)."""
    def __init__(self, parent, batch_folder_path=None, user_data=None):
        super().__init__(parent)
        self.title("Product Testing & Visual Inspection")
        
        # Maximize / Full window size setup
        try:
            self.state("zoomed")  # For Windows
        except Exception:
            self.attributes("-zoomed", True)  # For Linux/Mac fallback
            
        self.transient(parent)
        self.grab_set()

        self.batch_folder_path = batch_folder_path

        # Logged-in user passed from MainApp.  This is the authoritative
        # identity used when saving Product Testing / Visual Inspection data.
        self.user_data = user_data if isinstance(user_data, dict) else {}

        self.product_id_var = tk.StringVar()
        self.front_photo_path = None
        self.back_photo_path = None

        # Standard checklist items
        self.checks = {
            "No Scratches / Surface Defects": tk.BooleanVar(value=False),
            "Dimensions & Fit Verified": tk.BooleanVar(value=False),
            "Secure Components / Assembly": tk.BooleanVar(value=False),
            "Component Alignment Checked": tk.BooleanVar(value=False),
            "No Solder Bridges": tk.BooleanVar(value=False),
            "Polarity Orientation Correct": tk.BooleanVar(value=False)
        }

        # Dynamic conditional tracking variables for Visual Inspection
        self.manual_soldering_var = tk.BooleanVar(value=False)
        self.component_damage_var = tk.BooleanVar(value=False)
        self.developer_note_var = tk.BooleanVar(value=False)

        # Auto Testing Unit variables
        self.controller_prog_var = tk.BooleanVar(value=False)
        self.wifi_prog_var = tk.BooleanVar(value=False)
        self.tested_ok_var = tk.BooleanVar(value=False)

        # Repair variable
        self.repair_required_var = tk.BooleanVar(value=False)

        self.create_widgets()

    def create_widgets(self):
        # Configure root window weights to ensure proper resizing
        self.rowconfigure(0, weight=1)
        self.columnconfigure(0, weight=1)

        # Canvas with a Scrollbar for smooth navigation across resolutions
        canvas = tk.Canvas(self, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        
        main_frame = ttk.Frame(canvas, padding=20)
        main_frame.columnconfigure(0, weight=1)

        main_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas_window = canvas.create_window((0, 0), window=main_frame, anchor="nw")
        
        # Dynamically resize the canvas window width to match the viewport width
        def on_canvas_configure(event):
            canvas.itemconfig(canvas_window, width=event.width)
        canvas.bind("<Configure>", on_canvas_configure)

        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        # 1. Product ID Section & Search (Spans full width at top)
        id_frame = ttk.LabelFrame(main_frame, text=" Product Identification & Search ", padding=12)
        id_frame.pack(fill="x", pady=(0, 15), padx=5)
        id_frame.columnconfigure(1, weight=1)

        ttk.Label(id_frame, text="Enter Product ID:", font=("Helvetica", 10, "bold")).grid(row=0, column=0, sticky="w", padx=5, pady=5)
        ttk.Entry(id_frame, textvariable=self.product_id_var, width=30).grid(row=0, column=1, sticky="ew", padx=5, pady=5)
        ttk.Button(id_frame, text="🔍 Search / Edit", command=self.search_product_record).grid(row=0, column=2, padx=10, pady=5)

        # Create a container frame for two-column alignment
        columns_frame = ttk.Frame(main_frame)
        columns_frame.pack(fill="both", expand=True, pady=(0, 15), padx=5)
        columns_frame.columnconfigure(0, weight=1)
        columns_frame.columnconfigure(1, weight=1)

        # --- LEFT COLUMN ---
        left_column = ttk.Frame(columns_frame)
        left_column.grid(row=0, column=0, sticky="nsew", padx=(0, 10))

        # 2. Visual Inspection Checklist
        chk_frame = ttk.LabelFrame(left_column, text=" Visual Inspection Checklist ", padding=12)
        chk_frame.pack(fill="both", expand=True, pady=(0, 15))

        for text, var in self.checks.items():
            ttk.Checkbutton(chk_frame, text=text, variable=var).pack(anchor="w", pady=4, padx=5)

        # Conditional Item 1: Manually Soldering Required
        soldering_frame = ttk.Frame(chk_frame)
        soldering_frame.pack(fill="x", pady=3, padx=5)
        ttk.Checkbutton(soldering_frame, text="Manually Soldering Required", variable=self.manual_soldering_var, command=self.toggle_soldering_entry).pack(anchor="w")
        
        self.soldering_entry_frame = ttk.Frame(chk_frame)
        ttk.Label(self.soldering_entry_frame, text="Component Name(s):", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.soldering_entry = ttk.Entry(self.soldering_entry_frame, width=40)
        self.soldering_entry.pack(fill="x", anchor="w", padx=20, pady=2)

        # Conditional Item 2: Any Component Damage
        damage_frame = ttk.Frame(chk_frame)
        damage_frame.pack(fill="x", pady=3, padx=5)
        ttk.Checkbutton(damage_frame, text="Any Component Damage", variable=self.component_damage_var, command=self.toggle_damage_entry).pack(anchor="w")

        self.damage_entry_frame = ttk.Frame(chk_frame)
        ttk.Label(self.damage_entry_frame, text="Damage Details & Component Info:", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.damage_entry = ttk.Entry(self.damage_entry_frame, width=40)
        self.damage_entry.pack(fill="x", anchor="w", padx=20, pady=2)

        # Conditional Item 3: Developer Note
        dev_note_frame = ttk.Frame(chk_frame)
        dev_note_frame.pack(fill="x", pady=3, padx=5)
        self.dev_note_chk = ttk.Checkbutton(dev_note_frame, text="Developer Note", variable=self.developer_note_var, command=self.toggle_developer_note_entry)
        self.dev_note_chk.pack(anchor="w")

        self.dev_note_entry_frame = ttk.Frame(chk_frame)
        ttk.Label(self.dev_note_entry_frame, text="Developer Notes / Comments:", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.dev_note_entry = ttk.Entry(self.dev_note_entry_frame, width=40)
        self.dev_note_entry.pack(fill="x", anchor="w", padx=20, pady=2)


        # --- RIGHT COLUMN ---
        right_column = ttk.Frame(columns_frame)
        right_column.grid(row=0, column=1, sticky="nsew", padx=(10, 0))

        # 4. Connect to Auto Testing Unit Section
        auto_frame = ttk.LabelFrame(right_column, text=" Connect to Auto Testing Unit ", padding=12)
        auto_frame.pack(fill="x", pady=(0, 15))

        # Controller Programming Complete
        self.controller_chk = ttk.Checkbutton(auto_frame, text="Controller Programming Complete", variable=self.controller_prog_var, command=self.toggle_controller_note)
        self.controller_chk.pack(anchor="w", pady=4, padx=5)

        self.controller_note_frame = ttk.Frame(auto_frame)
        ttk.Label(self.controller_note_frame, text="Fault Note (Controller):", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.controller_note_entry = ttk.Entry(self.controller_note_frame, width=40)
        self.controller_note_entry.pack(fill="x", anchor="w", padx=20, pady=2)

        # Wifi Programming Complete
        self.wifi_chk = ttk.Checkbutton(auto_frame, text="Wifi Programming Complete", variable=self.wifi_prog_var, command=self.toggle_wifi_note)
        self.wifi_chk.pack(anchor="w", pady=4, padx=5)

        self.wifi_note_frame = ttk.Frame(auto_frame)
        ttk.Label(self.wifi_note_frame, text="Fault Note (Wifi):", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.wifi_note_entry = ttk.Entry(self.wifi_note_frame, width=40)
        self.wifi_note_entry.pack(fill="x", anchor="w", padx=20, pady=2)

        # Tested OK
        self.tested_chk = ttk.Checkbutton(auto_frame, text="Tested OK", variable=self.tested_ok_var, command=self.toggle_tested_note)
        self.tested_chk.pack(anchor="w", pady=4, padx=5)

        self.tested_note_frame = ttk.Frame(auto_frame)
        ttk.Label(self.tested_note_frame, text="Fault Note (Testing):", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.tested_note_entry = ttk.Entry(self.tested_note_frame, width=40)
        self.tested_note_entry.pack(fill="x", anchor="w", padx=20, pady=2)

        # 5. Repair Status Section
        repair_frame = ttk.LabelFrame(right_column, text=" Repair Status ", padding=12)
        repair_frame.pack(fill="x", pady=(0, 15))

        self.repair_chk = ttk.Checkbutton(repair_frame, text="Repair Required / Repaired", variable=self.repair_required_var, command=self.toggle_repair_note)
        self.repair_chk.pack(anchor="w", pady=4, padx=5)

        self.repair_note_frame = ttk.Frame(repair_frame)
        ttk.Label(self.repair_note_frame, text="Repair / Action Notes:", font=("Helvetica", 9)).pack(anchor="w", padx=20, pady=(2, 0))
        self.repair_note_entry = ttk.Entry(self.repair_note_frame, width=40)
        self.repair_note_entry.pack(fill="x", anchor="w", padx=20, pady=2)

        # 3. Photo Upload Section
        photo_frame = ttk.LabelFrame(right_column, text=" Individual Product Photos ", padding=12)
        photo_frame.pack(fill="x", pady=(0, 15))
        photo_frame.columnconfigure(1, weight=1)

        ttk.Button(photo_frame, text="Update Front Photo", command=lambda: self.load_photo("front")).grid(row=0, column=0, padx=5, pady=6)
        ttk.Button(photo_frame, text="Camera Front", command=lambda: self.capture_photo("front")).grid(row=0, column=1, padx=5, pady=6)
        self.lbl_front = ttk.Label(photo_frame, text="No front photo selected", font=("Helvetica", 9, "italic"))
        self.lbl_front.grid(row=1, column=0, columnspan=2, sticky="w", padx=5, pady=(0, 8))

        ttk.Button(photo_frame, text="Update Back Photo", command=lambda: self.load_photo("back")).grid(row=2, column=0, padx=5, pady=6)
        ttk.Button(photo_frame, text="Camera Back", command=lambda: self.capture_photo("back")).grid(row=2, column=1, padx=5, pady=6)
        self.lbl_back = ttk.Label(photo_frame, text="No back photo selected", font=("Helvetica", 9, "italic"))
        self.lbl_back.grid(row=3, column=0, columnspan=2, sticky="w", padx=5, pady=(0, 4))

        # Save Button (Centered at the bottom)
        self.btn_save = ttk.Button(main_frame, text="💾 Save / Update Inspection Record", command=self.save_inspection_data)
        self.btn_save.pack(pady=15)

        # Initialize view states
        self.toggle_soldering_entry()
        self.toggle_damage_entry()
        self.toggle_developer_note_entry()
        self.toggle_controller_note()
        self.toggle_wifi_note()
        self.toggle_tested_note()
        self.toggle_repair_note()

    def toggle_soldering_entry(self):
        if self.manual_soldering_var.get():
            self.soldering_entry_frame.pack(fill="x", pady=2)
        else:
            self.soldering_entry_frame.pack_forget()
            self.soldering_entry.delete(0, tk.END)

    def toggle_damage_entry(self):
        if self.component_damage_var.get():
            self.damage_entry_frame.pack(fill="x", pady=2)
        else:
            self.damage_entry_frame.pack_forget()
            self.damage_entry.delete(0, tk.END)

    def toggle_developer_note_entry(self):
        if self.developer_note_var.get():
            self.dev_note_entry_frame.pack(fill="x", pady=2, after=self.dev_note_chk.master)
        else:
            self.dev_note_entry_frame.pack_forget()
            self.dev_note_entry.delete(0, tk.END)

    def toggle_controller_note(self):
        if not self.controller_prog_var.get():
            self.controller_note_frame.pack(after=self.controller_chk, fill="x", pady=2)
        else:
            self.controller_note_frame.pack_forget()
            self.controller_note_entry.delete(0, tk.END)

    def toggle_wifi_note(self):
        if not self.wifi_prog_var.get():
            self.wifi_note_frame.pack(after=self.wifi_chk, fill="x", pady=2)
        else:
            self.wifi_note_frame.pack_forget()
            self.wifi_note_entry.delete(0, tk.END)

    def toggle_tested_note(self):
        if not self.tested_ok_var.get():
            self.tested_note_frame.pack(after=self.tested_chk, fill="x", pady=2)
        else:
            self.tested_note_frame.pack_forget()
            self.tested_note_entry.delete(0, tk.END)

    def toggle_repair_note(self):
        if self.repair_required_var.get():
            self.repair_note_frame.pack(after=self.repair_chk, fill="x", pady=2)
        else:
            self.repair_note_frame.pack_forget()
            self.repair_note_entry.delete(0, tk.END)

    def load_photo(self, photo_type):
        file_path = filedialog.askopenfilename(
            title=f"Select {photo_type.capitalize()} View Photo",
            filetypes=[("Image Files", "*.jpg *.jpeg *.png *.bmp")],
            parent=self
        )
        if file_path:
            if photo_type == "front":
                self.front_photo_path = file_path
                self.lbl_front.config(text=os.path.basename(file_path))
            else:
                self.back_photo_path = file_path
                self.lbl_back.config(text=os.path.basename(file_path))

    def get_product_photo_folder(self):
        base_target_dir = self.batch_folder_path if self.batch_folder_path and os.path.exists(self.batch_folder_path) else os.path.join(os.getcwd(), "Production")
        prod_id = self.product_id_var.get().strip()
        if not prod_id:
            return None
        return os.path.join(base_target_dir, "Product_Testing_Records", prod_id)

    def capture_photo(self, photo_type):
        if not HAS_CV2:
            messagebox.showwarning(
                "Camera Not Available",
                "OpenCV is not installed. Please upload an image instead.",
                parent=self,
            )
            return

        capture_dir = self.get_product_photo_folder()
        if not capture_dir:
            messagebox.showwarning("Product ID Required", "Please enter the Product ID before capturing the photo.", parent=self)
            return
        os.makedirs(capture_dir, exist_ok=True)

        camera = cv2.VideoCapture(0)
        captured_path = ""
        preview_title = f"Capture {photo_type.capitalize()} Photo"

        try:
            if not camera.isOpened():
                raise RuntimeError("No camera detected or camera access was denied.")

            while True:
                ret, frame = camera.read()
                if not ret or frame is None:
                    raise RuntimeError("Unable to capture an image from the camera.")

                cv2.putText(
                    frame,
                    f"Press SPACE to capture {photo_type} photo | ESC to cancel",
                    (12, 30),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.7,
                    (0, 255, 0),
                    2,
                )
                cv2.imshow(preview_title, frame)

                key = cv2.waitKey(1) & 0xFF
                if key == 32:  # space
                    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                    captured_path = os.path.join(capture_dir, f"{photo_type}_view_{timestamp}.png")
                    cv2.imwrite(captured_path, frame)
                    break
                if key == 27:  # ESC
                    captured_path = ""
                    break

            if not captured_path:
                return

            if photo_type == "front":
                self.front_photo_path = captured_path
                self.lbl_front.config(text=os.path.basename(captured_path))
            else:
                self.back_photo_path = captured_path
                self.lbl_back.config(text=os.path.basename(captured_path))

            messagebox.showinfo("Photo Captured", f"{photo_type.capitalize()} photo saved successfully.\nSaved in: {capture_dir}", parent=self)
        except Exception as exc:
            messagebox.showerror("Camera Capture Failed", f"Could not capture the photo:\n{exc}", parent=self)
        finally:
            camera.release()
            cv2.destroyAllWindows()

    def search_product_record(self):
        prod_id = self.product_id_var.get().strip()
        if not prod_id:
            messagebox.showwarning("Validation Error", "Please enter a Product ID to search.", parent=self)
            return

        base_target_dir = self.batch_folder_path if self.batch_folder_path and os.path.exists(self.batch_folder_path) else os.path.join(os.getcwd(), "Production")
        csv_file = os.path.join(base_target_dir, "product_testing_log.csv")

        record_found = False
        if os.path.exists(csv_file):
            try:
                with open(csv_file, mode="r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    for row in reader:
                        if row.get("Product_ID") == prod_id:
                            record_found = True
                            for text, var in self.checks.items():
                                if text in row:
                                    var.set(row[text] == "Passed")
                            
                            soldering_val = row.get("Manual_Soldering_Required", "No")
                            if soldering_val.startswith("Yes:"):
                                self.manual_soldering_var.set(True)
                                self.soldering_entry.delete(0, tk.END)
                                self.soldering_entry.insert(0, soldering_val.split(":", 1)[1].strip())
                            else:
                                self.manual_soldering_var.set(False)
                            self.toggle_soldering_entry()

                            damage_val = row.get("Component_Damage_Details", "None")
                            if damage_val.startswith("Yes:"):
                                self.component_damage_var.set(True)
                                self.damage_entry.delete(0, tk.END)
                                self.damage_entry.insert(0, damage_val.split(":", 1)[1].strip())
                            else:
                                self.component_damage_var.set(False)
                            self.toggle_damage_entry()

                            dev_val = row.get("Developer_Note", "None")
                            if dev_val.startswith("Yes:"):
                                self.developer_note_var.set(True)
                                self.dev_note_entry.delete(0, tk.END)
                                self.dev_note_entry.insert(0, dev_val.split(":", 1)[1].strip())
                            else:
                                self.developer_note_var.set(False)
                            self.toggle_developer_note_entry()

                            c_prog = row.get("Controller_Programming", "")
                            self.controller_prog_var.set("Complete" in c_prog)
                            if "Note:" in c_prog:
                                self.controller_note_entry.delete(0, tk.END)
                                self.controller_note_entry.insert(0, c_prog.split("Note:", 1)[1].strip())
                            self.toggle_controller_note()

                            w_prog = row.get("Wifi_Programming", "")
                            self.wifi_prog_var.set("Complete" in w_prog)
                            if "Note:" in w_prog:
                                self.wifi_note_entry.delete(0, tk.END)
                                self.wifi_note_entry.insert(0, w_prog.split("Note:", 1)[1].strip())
                            self.toggle_wifi_note()

                            t_ok = row.get("Tested_OK", "")
                            self.tested_ok_var.set("Yes" in t_ok)
                            if "Note:" in t_ok:
                                self.tested_note_entry.delete(0, tk.END)
                                self.tested_note_entry.insert(0, t_ok.split("Note:", 1)[1].strip())
                            self.toggle_tested_note()

                            repair_val = row.get("Repair_Status", "No")
                            if repair_val.startswith("Yes:"):
                                self.repair_required_var.set(True)
                                self.repair_note_entry.delete(0, tk.END)
                                self.repair_note_entry.insert(0, repair_val.split(":", 1)[1].strip())
                            else:
                                self.repair_required_var.set(False)
                            self.toggle_repair_note()
                            
                            self.lbl_front.config(text=row.get("Front_Photo", "No front photo selected"))
                            self.lbl_back.config(text=row.get("Back_Photo", "No back photo selected"))
                            break
            except Exception as e:
                messagebox.showerror("Read Error", f"Failed to search record:\n{e}", parent=self)
                return

        if record_found:
            messagebox.showinfo("Record Found", f"Existing record for '{prod_id}' loaded.", parent=self)
        else:
            messagebox.showinfo("Not Found", f"No existing record found for '{prod_id}'.", parent=self)

    def get_current_user_name(self):
        """Return the actual user who is currently logged in.

        MainApp passes the authenticated user record into ProductionView and
        then into this window.  Therefore every new/update inspection record
        is saved against the logged-in user instead of always using Admin.
        """
        if isinstance(self.user_data, dict):
            full_name = str(self.user_data.get("full_name", "")).strip()
            username = str(self.user_data.get("username", "")).strip()
            if full_name:
                return full_name
            if username:
                return username

        # Backward-compatible fallback for older code that sets
        # auth_manager.current_user globally.
        current_user = getattr(auth_manager, "current_user", {})
        if isinstance(current_user, dict):
            full_name = str(current_user.get("full_name", "")).strip()
            username = str(current_user.get("username", "")).strip()
            if full_name:
                return full_name
            if username:
                return username
        elif hasattr(current_user, "username"):
            full_name = str(getattr(current_user, "full_name", "")).strip()
            username = str(getattr(current_user, "username", "")).strip()
            if full_name:
                return full_name
            if username:
                return username

        return "Admin"

    def save_inspection_data(self):
        prod_id = self.product_id_var.get().strip()
        if not prod_id:
            messagebox.showwarning("Validation Error", "Please enter a valid Product ID.", parent=self)
            return

        base_target_dir = self.batch_folder_path if self.batch_folder_path and os.path.exists(self.batch_folder_path) else os.path.join(os.getcwd(), "Production")
        prod_folder = os.path.join(base_target_dir, "Product_Testing_Records", prod_id)
        os.makedirs(prod_folder, exist_ok=True)

        saved_front = "None"
        saved_back = "None"

        try:
            if self.front_photo_path and os.path.exists(self.front_photo_path):
                ext = os.path.splitext(self.front_photo_path)[1]
                saved_front = f"front_view{ext}"
                dest_front = os.path.join(prod_folder, saved_front)
                if HAS_PIL:
                    Image.open(self.front_photo_path).save(dest_front)
                else:
                    shutil.copy(self.front_photo_path, dest_front)

            if self.back_photo_path and os.path.exists(self.back_photo_path):
                ext = os.path.splitext(self.back_photo_path)[1]
                saved_back = f"back_view{ext}"
                dest_back = os.path.join(prod_folder, saved_back)
                if HAS_PIL:
                    Image.open(self.back_photo_path).save(dest_back)
                else:
                    shutil.copy(self.back_photo_path, dest_back)

            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            action_user = self.get_current_user_name()

            # --- User Name placed FIRST so it becomes Column A ---
            record = {
                "User_Name": action_user,
                "Timestamp": now_str,
                "Product_ID": prod_id,
            }
            for text, var in self.checks.items():
                record[text] = "Passed" if var.get() else "Failed"
            
            record["Manual_Soldering_Required"] = f"Yes: {self.soldering_entry.get().strip()}" if self.manual_soldering_var.get() and self.soldering_entry.get().strip() else ("Yes" if self.manual_soldering_var.get() else "No")
            record["Component_Damage_Details"] = f"Yes: {self.damage_entry.get().strip()}" if self.component_damage_var.get() and self.damage_entry.get().strip() else ("Yes" if self.component_damage_var.get() else "None")
            record["Developer_Note"] = f"Yes: {self.dev_note_entry.get().strip()}" if self.developer_note_var.get() and self.dev_note_entry.get().strip() else ("Yes" if self.developer_note_var.get() else "None")

            if self.controller_prog_var.get():
                record["Controller_Programming"] = "Complete"
            else:
                note = self.controller_note_entry.get().strip()
                record["Controller_Programming"] = f"Incomplete - Note: {note}" if note else "Incomplete"

            if self.wifi_prog_var.get():
                record["Wifi_Programming"] = "Complete"
            else:
                note = self.wifi_note_entry.get().strip()
                record["Wifi_Programming"] = f"Incomplete - Note: {note}" if note else "Incomplete"

            if self.tested_ok_var.get():
                record["Tested_OK"] = "Yes"
            else:
                note = self.tested_note_entry.get().strip()
                record["Tested_OK"] = f"No - Note: {note}" if note else "No"

            if self.repair_required_var.get():
                r_note = self.repair_note_entry.get().strip()
                record["Repair_Status"] = f"Yes: {r_note}" if r_note else "Yes"
            else:
                record["Repair_Status"] = "No"

            record["Front_Photo"] = saved_front if saved_front != "None" else self.lbl_front.cget("text")
            record["Back_Photo"] = saved_back if saved_back != "None" else self.lbl_back.cget("text")
            record["Last_Edited_By"] = action_user

            csv_file = os.path.join(base_target_dir, "product_testing_log.csv")
            rows = []
            updated = False
            
            # Ensure "User_Name" stays at the very beginning of the fieldnames list
            fieldnames = ["User_Name"]
            for key in record.keys():
                if key not in fieldnames:
                    fieldnames.append(key)
            
            if os.path.exists(csv_file):
                with open(csv_file, mode="r", newline="", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    for fld in (reader.fieldnames or []):
                        if fld not in fieldnames:
                            fieldnames.append(fld)
                    for row in reader:
                        if row.get("Product_ID") == prod_id:
                            row.update(record)
                            updated = True
                        rows.append(row)

            if not updated:
                rows.append(record)

            with open(csv_file, mode="w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)

            # --- Excel Integration Block ---
            if HAS_OPENPYXL:
                excel_file = os.path.join(base_target_dir, "Product_Testing_Report.xlsx")
                if os.path.exists(excel_file):
                    wb = openpyxl.load_workbook(excel_file)
                    ws = wb.active
                else:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Testing Logs"
                    ws.append(fieldnames)

                excel_headers = [cell.value for cell in ws[1] if cell.value is not None]
                
                # Ensure all fields exist in the headers
                for fld in fieldnames:
                    if fld not in excel_headers:
                        excel_headers.append(fld)

                # Force ensure "User_Name" is at Column A (Index 1)
                if "User_Name" in excel_headers:
                    excel_headers.remove("User_Name")
                excel_headers.insert(0, "User_Name")

                # Re-write or align header row in Excel to match new structure
                for idx, h in enumerate(excel_headers, start=1):
                    ws.cell(row=1, column=idx, value=h)

                prod_id_col_idx = excel_headers.index("Product_ID") + 1 if "Product_ID" in excel_headers else 3
                row_idx_to_update = None
                
                for row_idx in range(2, ws.max_row + 1):
                    if str(ws.cell(row=row_idx, column=prod_id_col_idx).value) == prod_id:
                        row_idx_to_update = row_idx
                        break

                if row_idx_to_update:
                    for key, val in record.items():
                        if key in excel_headers:
                            col_idx = excel_headers.index(key) + 1
                            ws.cell(row=row_idx_to_update, column=col_idx, value=val)
                else:
                    new_row = [record.get(h, "") for h in excel_headers]
                    ws.append(new_row)

                wb.save(excel_file)

            messagebox.showinfo("Success", f"Testing record for Product ID '{prod_id}' saved successfully!", parent=self)
            self.destroy()

        except Exception as e:
            messagebox.showerror("Save Error", f"Failed to save testing data:\n{e}", parent=self)
            
class ProductionView(ttk.Frame):
    """Production Module updated with Product Testing integrated inside Existing Production Directory."""

    def __init__(self, parent, user_data=None):
        super().__init__(parent)

        # Keep the authenticated user available to all Production screens.
        self.user_data = user_data if isinstance(user_data, dict) else {}

        self.prod_dir = os.path.join(os.getcwd(), "Production")
        self.new_prod_dir = os.path.join(self.prod_dir, "New Production Start")
        self.csv_path = os.path.join(self.prod_dir, "products_list.csv")

        self.products = []
        self.attached_file_path = None
        self.selected_personnel = {}
        self.current_materials = []
        self.batch_map = {}
        
        # Initialized to False so sections 2, 3, and 4 are folded by default
        self.is_mat_frame_expanded = False
        self.is_pcb_frame_expanded = False
        self.is_receive_frame_expanded = False

        self.ensure_production_storage()
        self.load_products_from_csv()
        self.show_mode_selection()

    def ensure_production_storage(self):
        os.makedirs(self.prod_dir, exist_ok=True)
        os.makedirs(self.new_prod_dir, exist_ok=True)

        if not os.path.exists(self.csv_path):
            with open(
                self.csv_path, mode="w", newline="", encoding="utf-8"
            ) as f:
                writer = csv.writer(f)
                writer.writerow(["Product_Name", "Created_Date"])

    def load_products_from_csv(self):
        self.products = []
        if os.path.exists(self.csv_path):
            try:
                with open(self.csv_path, mode="r", encoding="utf-8") as f:
                    reader = csv.reader(f)
                    next(reader, None)
                    for row in reader:
                        if row and row[0].strip():
                            p_name = row[0].strip()
                            if p_name not in self.products:
                                self.products.append(p_name)
            except Exception as e:
                print(f"Error reading product CSV: {e}")

    def save_product_to_csv(self, product_name):
        if product_name in self.products:
            return

        self.products.append(product_name)
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        try:
            with open(
                self.csv_path, mode="a", newline="", encoding="utf-8"
            ) as f:
                writer = csv.writer(f)
                writer.writerow([product_name, now_str])
        except Exception as e:
            messagebox.showerror(
                "Storage Error", f"Failed to save product:\n{e}", parent=self
            )

    def resequence_sr_no(self):
        for idx, item in enumerate(self.current_materials, start=1):
            item["sr_no"] = str(idx)

    def clear_view(self):
        for child in self.winfo_children():
            child.destroy()

    def show_mode_selection(self):
        self.clear_view()

        container = ttk.Frame(self, padding=40)
        container.pack(expand=True)

        ttk.Label(
            container, text="Production Management", font=("Helvetica", 16, "bold")
        ).pack(pady=(0, 20))

        btn_frame = ttk.Frame(container)
        btn_frame.pack()

        ttk.Button(
            btn_frame,
            text="➕ New Production Start",
            width=32,
            command=self.setup_new_production_ui,
        ).pack(pady=10)
        ttk.Button(
            btn_frame,
            text="📁 Existing Production Directory",
            width=32,
            command=self.show_existing_production_ui,
        ).pack(pady=10)
        ttk.Button(
            btn_frame,
            text="production Details",
            width=32,
            command=self.show_production_details,
        ).pack(pady=10)

    # ------------------------------------------------------------------
    # PRODUCTION DETAILS
    # ------------------------------------------------------------------

    def show_production_details(self):
        self.clear_view()

        top_bar = ttk.Frame(self, padding=10)
        top_bar.pack(fill="x")

        ttk.Button(
            top_bar, text="⬅ Back to Option", command=self.show_mode_selection
        ).pack(side="left")
        ttk.Label(
            top_bar,
            text="production Details",
            font=("Helvetica", 13, "bold"),
        ).pack(side="left", padx=15)

        # Main scrollable container
        main_canvas = tk.Canvas(self, borderwidth=0, highlightthickness=0)
        scrollbar = ttk.Scrollbar(
            self, orient="vertical", command=main_canvas.yview
        )
        scrollable_frame = ttk.Frame(main_canvas, padding=20)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: main_canvas.configure(
                scrollregion=main_canvas.bbox("all")
            ),
        )
        main_canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        main_canvas.configure(yscrollcommand=scrollbar.set)

        main_canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # Batch Selection Section
        batch_frame = ttk.LabelFrame(
            scrollable_frame, text=" Select Production Batch ", padding=15
        )
        batch_frame.pack(fill="x", pady=(0, 15))

        ttk.Label(
            batch_frame, text="Select Batch:", font=("Helvetica", 10, "bold")
        ).grid(row=0, column=0, sticky="w", padx=5, pady=5)

        self.details_combo_batch = ttk.Combobox(batch_frame, width=45, state="readonly")
        self.details_combo_batch.grid(row=0, column=1, sticky="w", padx=10, pady=5)
        self.details_combo_batch.bind(
            "<<ComboboxSelected>>", self.on_details_batch_selected
        )

        self.details_batch_info = ttk.Label(
            batch_frame,
            text="Select a batch from the dropdown above to view details.",
            font=("Helvetica", 9, "italic"),
        )
        self.details_batch_info.grid(
            row=1, column=0, columnspan=3, sticky="w", padx=5, pady=5
        )

        # Batch Details Container (initially hidden)
        self.details_container = ttk.Frame(scrollable_frame)
        
        # Load batches into dropdown
        self.load_batches_into_details_dropdown()

    def load_batches_into_details_dropdown(self):
        self.details_batch_map = {}
        batch_options = []

        pattern = os.path.join(self.new_prod_dir, "*", "info.csv")
        csv_files = glob.glob(pattern)

        for csv_file in csv_files:
            try:
                with open(csv_file, mode="r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    row = next(reader, None)
                    if row:
                        p_name = row.get("Product_Name", "Unknown")
                        b_num = row.get("Batch_Number", "Unknown")
                        display_str = f"{b_num} ({p_name})"

                        self.details_batch_map[display_str] = (csv_file, row)
                        batch_options.append(display_str)
            except Exception as e:
                print(f"Error reading {csv_file}: {e}")

        self.details_combo_batch["values"] = batch_options
        if batch_options:
            self.details_combo_batch.set("")

    def on_details_batch_selected(self, event):
        selected_key = self.details_combo_batch.get()
        if not selected_key or selected_key not in self.details_batch_map:
            return

        # Clear previous details
        for widget in self.details_container.winfo_children():
            widget.destroy()
        self.details_container.pack_forget()

        csv_path, row_data = self.details_batch_map[selected_key]
        
        # Show details container
        self.details_container.pack(fill="both", expand=True, pady=(10, 0))

        # Basic Information Section
        info_frame = ttk.LabelFrame(
            self.details_container, text=" Basic Batch Information ", padding=15
        )
        info_frame.pack(fill="x", pady=(0, 15))

        info_grid = ttk.Frame(info_frame)
        info_grid.pack(fill="x")

        # Display all key information from the batch
        info_items = [
            ("Product Name:", row_data.get("Product_Name", "N/A")),
            ("Batch Number:", row_data.get("Batch_Number", "N/A")),
            ("Start Date:", row_data.get("Start_Date", "N/A")),
            ("Expected Completion Date:", row_data.get("Expected_Completion_Date", "N/A")),
            ("Assigned Personnel:", row_data.get("Assigned_Personnel", "N/A")),
        ]

        for idx, (label, value) in enumerate(info_items):
            ttk.Label(info_grid, text=label, font=("Helvetica", 10, "bold")).grid(
                row=idx // 2, column=(idx % 2) * 2, sticky="w", padx=5, pady=5
            )
            ttk.Label(info_grid, text=value, font=("Helvetica", 10)).grid(
                row=idx // 2, column=(idx % 2) * 2 + 1, sticky="w", padx=5, pady=5
            )

        # Product Testing & Visual Inspection Report Section
        testing_frame = ttk.LabelFrame(
            self.details_container, text=" Product Testing & Visual Inspection Report ", padding=15
        )
        testing_frame.pack(fill="both", expand=True, pady=(0, 15))

        # Download Excel Button
        download_btn_frame = ttk.Frame(testing_frame)
        download_btn_frame.pack(fill="x", pady=(0, 10))
        
        ttk.Button(
            download_btn_frame,
            text="📥 Download Excel Report",
            command=lambda: self.download_testing_excel(csv_path),
        ).pack(side="right")

        # Testing data treeview
        test_cols = (
            "Product_ID",
            "User_Name",
            "Timestamp",
            "Visual_Inspection",
            "Controller_Prog",
            "Wifi_Prog",
            "Tested_OK",
            "Repair_Status",
        )
        details_test_tree = ttk.Treeview(
            testing_frame,
            columns=test_cols,
            show="headings",
            selectmode="browse",
            height=8,
        )

        details_test_tree.heading("Product_ID", text="Product ID")
        details_test_tree.heading("User_Name", text="User Name")
        details_test_tree.heading("Timestamp", text="Timestamp")
        details_test_tree.heading("Visual_Inspection", text="Visual Inspection")
        details_test_tree.heading("Controller_Prog", text="Controller Prog")
        details_test_tree.heading("Wifi_Prog", text="Wifi Prog")
        details_test_tree.heading("Tested_OK", text="Tested OK")
        details_test_tree.heading("Repair_Status", text="Repair Status")

        details_test_tree.column("Product_ID", width=100, anchor="w")
        details_test_tree.column("User_Name", width=120, anchor="w")
        details_test_tree.column("Timestamp", width=150, anchor="w")
        details_test_tree.column("Visual_Inspection", width=120, anchor="center")
        details_test_tree.column("Controller_Prog", width=100, anchor="center")
        details_test_tree.column("Wifi_Prog", width=80, anchor="center")
        details_test_tree.column("Tested_OK", width=80, anchor="center")
        details_test_tree.column("Repair_Status", width=100, anchor="center")

        vsb_test = ttk.Scrollbar(testing_frame, orient="vertical", command=details_test_tree.yview)
        details_test_tree.configure(yscrollcommand=vsb_test.set)
        
        details_test_tree.pack(side="left", fill="both", expand=True)
        vsb_test.pack(side="right", fill="y")

        # Load testing data for this batch
        self.load_testing_data_for_batch(details_test_tree, row_data)

        # Additional Information Section
        additional_frame = ttk.LabelFrame(
            self.details_container, text=" Additional Information ", padding=15
        )
        additional_frame.pack(fill="x", pady=(0, 15))

        ttk.Label(
            additional_frame,
            text=f"Batch Directory: {os.path.dirname(csv_path)}",
            font=("Helvetica", 9)
        ).pack(anchor="w", pady=2)
        
        # Show testing records count
        testing_log_path = os.path.join(os.path.dirname(csv_path), "product_testing_log.csv")
        if os.path.exists(testing_log_path):
            try:
                with open(testing_log_path, mode="r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    record_count = sum(1 for row in reader)
                ttk.Label(
                    additional_frame,
                    text=f"Testing Records: {record_count}",
                    font=("Helvetica", 9)
                ).pack(anchor="w", pady=2)
            except:
                ttk.Label(
                    additional_frame,
                    text="Testing Records: Unable to count",
                    font=("Helvetica", 9)
                ).pack(anchor="w", pady=2)
        else:
            ttk.Label(
                additional_frame,
                text="Testing Records: No records found",
                font=("Helvetica", 9)
            ).pack(anchor="w", pady=2)

        # Update the info label
        exp_date = row_data.get("Expected_Completion_Date", "N/A")
        start_date = row_data.get("Start_Date", "N/A")
        personnel = row_data.get("Assigned_Personnel", "None")

        self.details_batch_info.config(
            text=f"Start Date: {start_date}   |   Expected Completion: {exp_date}   |   Personnel: {personnel}",
            font=("Helvetica", 9, "normal"),
        )

    def load_testing_data_for_batch(self, tree_widget, batch_row_data):
        """Load product testing records for the selected batch."""
        tree_widget.delete(*tree_widget.get_children())
        
        # Get the batch directory
        batch_dir = os.path.dirname(self.details_batch_map[self.details_combo_batch.get()][0])
        
        # Look for testing log CSV
        testing_log_path = os.path.join(batch_dir, "product_testing_log.csv")
        
        if not os.path.exists(testing_log_path):
            tree_widget.insert("", "end", values=("No testing records found", "", "", "", "", "", "", ""))
            return
        
        try:
            with open(testing_log_path, mode="r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for row in reader:
                    # Calculate visual inspection status
                    visual_status = "Passed"
                    checklist_items = [
                        "No Scratches / Surface Defects",
                        "Dimensions & Fit Verified", 
                        "Secure Components / Assembly",
                        "Component Alignment Checked",
                        "No Solder Bridges",
                        "Polarity Orientation Correct"
                    ]
                    for item in checklist_items:
                        if row.get(item, "Failed") == "Failed":
                            visual_status = "Failed"
                            break
                    
                    # Get controller programming status
                    controller_status = "Complete" if "Complete" in row.get("Controller_Programming", "") else "Incomplete"
                    
                    # Get wifi programming status
                    wifi_status = "Complete" if "Complete" in row.get("Wifi_Programming", "") else "Incomplete"
                    
                    # Get tested status
                    tested_status = "Yes" if "Yes" in row.get("Tested_OK", "") else "No"
                    
                    # Get repair status
                    repair_status = "Yes" if row.get("Repair_Status", "No").startswith("Yes") else "No"
                    
                    tree_widget.insert("", "end", values=(
                        row.get("Product_ID", ""),
                        row.get("User_Name", ""),
                        row.get("Timestamp", ""),
                        visual_status,
                        controller_status,
                        wifi_status,
                        tested_status,
                        repair_status,
                    ))
        except Exception as e:
            tree_widget.insert("", "end", values=(f"Error loading data: {str(e)}", "", "", "", "", "", "", ""))

    def download_testing_excel(self, batch_csv_path):
        """Download the testing report as Excel file."""
        if not HAS_OPENPYXL:
            messagebox.showwarning(
                "Excel Export Not Available",
                "openpyxl library is not installed. Please install it to export Excel files.",
                parent=self
            )
            return
        
        batch_dir = os.path.dirname(batch_csv_path)
        testing_log_path = os.path.join(batch_dir, "product_testing_log.csv")
        
        if not os.path.exists(testing_log_path):
            messagebox.showwarning(
                "No Testing Data",
                "No testing records found for this batch.",
                parent=self
            )
            return
        
        try:
            # Ask for save location
            save_path = filedialog.asksaveasfilename(
                title="Save Testing Report As",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
                initialfile=f"Testing_Report_{os.path.basename(batch_dir)}.xlsx",
                parent=self
            )
            
            if not save_path:
                return
            
            # Read the CSV data
            with open(testing_log_path, mode="r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                rows = list(reader)
            
            if not rows:
                messagebox.showwarning(
                    "No Data",
                    "No testing records found to export.",
                    parent=self
                )
                return
            
            # Create Excel workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Testing Report"
            
            # Define headers
            headers = [
                "User_Name",
                "Timestamp", 
                "Product_ID",
                "No Scratches / Surface Defects",
                "Dimensions & Fit Verified",
                "Secure Components / Assembly",
                "Component Alignment Checked",
                "No Solder Bridges",
                "Polarity Orientation Correct",
                "Manual_Soldering_Required",
                "Component_Damage_Details",
                "Developer_Note",
                "Controller_Programming",
                "Wifi_Programming",
                "Tested_OK",
                "Repair_Status",
                "Front_Photo",
                "Back_Photo",
                "Last_Edited_By"
            ]
            
            # Write headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
            
            # Write data
            for row_num, row_data in enumerate(rows, 2):
                for col_num, header in enumerate(headers, 1):
                    value = row_data.get(header, "")
                    ws.cell(row=row_num, column=col_num, value=value)
            
            # Auto-adjust column widths
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column].width = adjusted_width
            
            # Save the workbook
            wb.save(save_path)
            
            messagebox.showinfo(
                "Export Successful",
                f"Testing report exported successfully to:\n{save_path}",
                parent=self
            )
            
        except Exception as e:
            messagebox.showerror(
                "Export Error",
                f"Failed to export testing report:\n{e}",
                parent=self
            )

    # ------------------------------------------------------------------
    # EXISTING PRODUCTION DIRECTORY
    # ------------------------------------------------------------------

    def show_existing_production_ui(self):
        self.clear_view()

        top_bar = ttk.Frame(self, padding=10)
        top_bar.pack(fill="x")

        ttk.Button(
            top_bar, text="⬅ Back to Option", command=self.show_mode_selection
        ).pack(side="left")
        ttk.Label(
            top_bar,
            text="Existing Production Directory",
            font=("Helvetica", 13, "bold"),
        ).pack(side="left", padx=15)

        main_canvas = tk.Canvas(self, borderwidth=0, highlightthickness=0)
        scrollbar = ttk.Scrollbar(
            self, orient="vertical", command=main_canvas.yview
        )
        scrollable_frame = ttk.Frame(main_canvas, padding=15)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: main_canvas.configure(
                scrollregion=main_canvas.bbox("all")
            ),
        )
        main_canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        main_canvas.configure(yscrollcommand=scrollbar.set)

        main_canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # 1. Select Production Batch
        drop_frame = ttk.LabelFrame(
            scrollable_frame, text=" 1. Select Production Batch ", padding=15
        )
        drop_frame.pack(fill="x", pady=(0, 15))

        ttk.Label(
            drop_frame, text="Select Batch:", font=("Helvetica", 10, "bold")
        ).grid(row=0, column=0, sticky="w", padx=5, pady=5)

        self.combo_batch = ttk.Combobox(drop_frame, width=45, state="readonly")
        self.combo_batch.grid(row=0, column=1, sticky="w", padx=10, pady=5)
        self.combo_batch.bind(
            "<<ComboboxSelected>>", self.on_batch_dropdown_selected
        )

        self.lbl_batch_info = ttk.Label(
            drop_frame,
            text="Select a batch from the dropdown above.",
            font=("Helvetica", 9, "italic"),
        )
        self.lbl_batch_info.grid(
            row=1, column=0, columnspan=3, sticky="w", padx=5, pady=5
        )

        # 2. Material List & Purchase Order Frame (Collapsible)
        self.mat_frame = ttk.LabelFrame(
            scrollable_frame, text=" 2. Material List & Purchase Orders ", padding=10
        )
        self.mat_frame.pack(fill="both", expand=True, pady=(0, 15))

        mat_header_bar = ttk.Frame(self.mat_frame)
        mat_header_bar.pack(fill="x", side="top", pady=(0, 5))

        self.btn_toggle_mat = ttk.Button(
            mat_header_bar,
            text="▲ Fold Section",
            width=18,
            command=self.toggle_material_section,
        )
        self.btn_toggle_mat.pack(side="right")

        self.mat_content_container = ttk.Frame(self.mat_frame)
        self.mat_content_container.pack(fill="both", expand=True)

        if not self.is_mat_frame_expanded:
            self.mat_content_container.pack_forget()
            self.btn_toggle_mat.config(text="▼ Expand Section")

        mat_cols = (
            "PO_Select",
            "Sr_No",
            "Category",
            "Part_Name",
            "Value",
            "Package",
            "Tolerance",
            "Quantity",
            "QTY",
            "Rate",
            "Total_Price",
            "Make",
        )
        self.mat_tree = ttk.Treeview(
            self.mat_content_container,
            columns=mat_cols,
            show="headings",
            selectmode="browse",
            height=8,
        )

        self.mat_tree.tag_configure("evenrow", background="#ffffff")
        self.mat_tree.tag_configure("oddrow", background="#f2f2f2")

        self.mat_tree.heading("PO_Select", text="[ PO ]")
        self.mat_tree.heading("Sr_No", text="Sr. No.")
        self.mat_tree.heading("Category", text="Category")
        self.mat_tree.heading("Part_Name", text="Part Name")
        self.mat_tree.heading("Value", text="Value")
        self.mat_tree.heading("Package", text="Package")
        self.mat_tree.heading("Tolerance", text="Tolerance")
        self.mat_tree.heading("Quantity", text="Quantity")
        self.mat_tree.heading("QTY", text="PO QTY")
        self.mat_tree.heading("Rate", text="RATE")
        self.mat_tree.heading("Total_Price", text="Total Price")
        self.mat_tree.heading("Make", text="MAKE")

        self.mat_tree.column("PO_Select", width=50, anchor="center")
        self.mat_tree.column("Sr_No", width=55, anchor="center")
        self.mat_tree.column("Category", width=120, anchor="w")
        self.mat_tree.column("Part_Name", width=130, anchor="w")
        self.mat_tree.column("Value", width=80, anchor="center")
        self.mat_tree.column("Package", width=80, anchor="center")
        self.mat_tree.column("Tolerance", width=70, anchor="center")
        self.mat_tree.column("Quantity", width=65, anchor="center")
        self.mat_tree.column("QTY", width=75, anchor="center")
        self.mat_tree.column("Rate", width=60, anchor="center")
        self.mat_tree.column("Total_Price", width=90, anchor="center")
        self.mat_tree.column("Make", width=80, anchor="center")

        vsb_bot = ttk.Scrollbar(
            self.mat_content_container, orient="vertical", command=self.mat_tree.yview
        )
        hsb_bot = ttk.Scrollbar(
            self.mat_content_container, orient="horizontal", command=self.mat_tree.xview
        )
        self.mat_tree.configure(
            yscrollcommand=vsb_bot.set, xscrollcommand=hsb_bot.set
        )

        self.mat_tree.pack(side="top", fill="both", expand=True)
        vsb_bot.pack(side="right", fill="y")
        hsb_bot.pack(side="bottom", fill="x")

        self.mat_tree.bind("<Button-1>", self.on_material_tree_click)

        mat_btn_bar = ttk.Frame(self.mat_content_container, padding=(0, 10, 0, 0))
        mat_btn_bar.pack(fill="x", side="bottom")

        ttk.Button(
            mat_btn_bar,
            text="☑ Select / Unmark All PO",
            command=self.toggle_select_all_po,
        ).pack(side="left", padx=4)
        ttk.Button(
            mat_btn_bar,
            text="📊 Import Excel Sheet (.xlsx)",
            command=self.import_and_convert_excel,
        ).pack(side="left", padx=4)
        ttk.Button(
            mat_btn_bar,
            text="➕ Add Manual Row",
            command=self.open_add_material_dialog,
        ).pack(side="left", padx=4)
        ttk.Button(
            mat_btn_bar,
            text="❌ Delete Row",
            command=self.remove_selected_material,
        ).pack(side="left", padx=4)
        ttk.Button(
            mat_btn_bar,
            text="📦 Generate Purchase Order (.xlsx)",
            command=self.send_purchase_order,
        ).pack(side="left", padx=4)
        ttk.Button(
            mat_btn_bar,
            text="💾 Save to CSV",
            command=self.save_material_list_changes,
        ).pack(side="left", padx=4)

        # 3. Deliver to PCB Mounting Frame
        self.pcb_frame = ttk.LabelFrame(
            scrollable_frame, text=" 3. Deliver to PCB Mounting ", padding=10
        )
        self.pcb_frame.pack(fill="both", expand=True, pady=(0, 15))

        pcb_header_bar = ttk.Frame(self.pcb_frame)
        pcb_header_bar.pack(fill="x", side="top", pady=(0, 5))

        self.btn_toggle_pcb = ttk.Button(
            pcb_header_bar,
            text="▲ Fold Section",
            width=18,
            command=self.toggle_pcb_section,
        )
        self.btn_toggle_pcb.pack(side="right")

        self.pcb_content_container = ttk.Frame(self.pcb_frame)
        self.pcb_content_container.pack(fill="both", expand=True)

        if not self.is_pcb_frame_expanded:
            self.pcb_content_container.pack_forget()
            self.btn_toggle_pcb.config(text="▼ Expand Section")

        pcb_cols = (
            "PCB_Select",
            "Sr_No",
            "Category",
            "Part_Name",
            "Value",
            "Package",
            "Tolerance",
            "Quantity",
            "QTY",
            "Make",
        )
        self.pcb_tree = ttk.Treeview(
            self.pcb_content_container,
            columns=pcb_cols,
            show="headings",
            selectmode="browse",
            height=8,
        )

        self.pcb_tree.tag_configure("evenrow", background="#ffffff")
        self.pcb_tree.tag_configure("oddrow", background="#f2f2f2")

        self.pcb_tree.heading("PCB_Select", text="[ Sel ]")
        self.pcb_tree.heading("Sr_No", text="Sr. No.")
        self.pcb_tree.heading("Category", text="Category")
        self.pcb_tree.heading("Part_Name", text="Part Name")
        self.pcb_tree.heading("Value", text="Value")
        self.pcb_tree.heading("Package", text="Package")
        self.pcb_tree.heading("Tolerance", text="Tolerance")
        self.pcb_tree.heading("Quantity", text="Quantity")
        self.pcb_tree.heading("QTY", text="PO QTY")
        self.pcb_tree.heading("Make", text="MAKE")

        self.pcb_tree.column("PCB_Select", width=50, anchor="center")
        self.pcb_tree.column("Sr_No", width=55, anchor="center")
        self.pcb_tree.column("Category", width=120, anchor="w")
        self.pcb_tree.column("Part_Name", width=130, anchor="w")
        self.pcb_tree.column("Value", width=80, anchor="center")
        self.pcb_tree.column("Package", width=80, anchor="center")
        self.pcb_tree.column("Tolerance", width=70, anchor="center")
        self.pcb_tree.column("Quantity", width=65, anchor="center")
        self.pcb_tree.column("QTY", width=75, anchor="center")
        self.pcb_tree.column("Make", width=80, anchor="center")

        vsb_pcb = ttk.Scrollbar(
            self.pcb_content_container, orient="vertical", command=self.pcb_tree.yview
        )
        hsb_pcb = ttk.Scrollbar(
            self.pcb_content_container, orient="horizontal", command=self.pcb_tree.xview
        )
        self.pcb_tree.configure(
            yscrollcommand=vsb_pcb.set, xscrollcommand=hsb_pcb.set
        )

        self.pcb_tree.pack(side="top", fill="both", expand=True)
        vsb_pcb.pack(side="right", fill="y")
        hsb_pcb.pack(side="bottom", fill="x")

        self.pcb_tree.bind("<Button-1>", self.on_pcb_tree_click)

        pcb_btn_bar = ttk.Frame(self.pcb_content_container, padding=(0, 10, 0, 0))
        pcb_btn_bar.pack(fill="x", side="bottom")

        ttk.Button(
            pcb_btn_bar,
            text="☑ Select / Unmark All",
            command=self.toggle_select_all_pcb,
        ).pack(side="left", padx=4)
        ttk.Button(
            pcb_btn_bar,
            text="📊 Import Excel Sheet (.xlsx)",
            command=self.import_and_convert_excel,
        ).pack(side="left", padx=4)
        ttk.Button(
            pcb_btn_bar,
            text="➕ Add Manual Row",
            command=self.open_add_material_dialog,
        ).pack(side="left", padx=4)
        ttk.Button(
            pcb_btn_bar,
            text="❌ Delete Row",
            command=self.remove_selected_material,
        ).pack(side="left", padx=4)
        ttk.Button(
            pcb_btn_bar,
            text="📦 Generate Material List (.xlsx)",
            command=self.send_pcb_material_list,
        ).pack(side="left", padx=4)
        ttk.Button(
            pcb_btn_bar,
            text="💾 Save to CSV",
            command=self.save_material_list_changes,
        ).pack(side="left", padx=4)

        # 4. Receive from PCB Mounting Frame
        self.receive_frame = ttk.LabelFrame(
            scrollable_frame, text=" 4. Receive from PCB Mounting ", padding=10
        )
        self.receive_frame.pack(fill="both", expand=True, pady=(0, 15))

        receive_header_bar = ttk.Frame(self.receive_frame)
        receive_header_bar.pack(fill="x", side="top", pady=(0, 5))

        self.btn_toggle_receive = ttk.Button(
            receive_header_bar,
            text="▲ Fold Section",
            width=18,
            command=self.toggle_receive_section,
        )
        self.btn_toggle_receive.pack(side="right")

        self.receive_content_container = ttk.Frame(self.receive_frame)
        self.receive_content_container.pack(fill="both", expand=True)

        if not self.is_receive_frame_expanded:
            self.receive_content_container.pack_forget()
            self.btn_toggle_receive.config(text="▼ Expand Section")

        receive_cols = (
            "Receive_Select",
            "Sr_No",
            "Category",
            "Part_Name",
            "Value",
            "Package",
            "Tolerance",
            "Quantity",
            "QTY",
            "Make",
        )
        self.receive_tree = ttk.Treeview(
            self.receive_content_container,
            columns=receive_cols,
            show="headings",
            selectmode="browse",
            height=8,
        )

        self.receive_tree.tag_configure("evenrow", background="#ffffff")
        self.receive_tree.tag_configure("oddrow", background="#f2f2f2")

        self.receive_tree.heading("Receive_Select", text="[ Rcv ]")
        self.receive_tree.heading("Sr_No", text="Sr. No.")
        self.receive_tree.heading("Category", text="Category")
        self.receive_tree.heading("Part_Name", text="Part Name")
        self.receive_tree.heading("Value", text="Value")
        self.receive_tree.heading("Package", text="Package")
        self.receive_tree.heading("Tolerance", text="Tolerance")
        self.receive_tree.heading("Quantity", text="Quantity")
        self.receive_tree.heading("QTY", text="PO QTY")
        self.receive_tree.heading("Make", text="Make")

        self.receive_tree.column("Receive_Select", width=50, anchor="center")
        self.receive_tree.column("Sr_No", width=55, anchor="center")
        self.receive_tree.column("Category", width=120, anchor="w")
        self.receive_tree.column("Part_Name", width=130, anchor="w")
        self.receive_tree.column("Value", width=80, anchor="center")
        self.receive_tree.column("Package", width=80, anchor="center")
        self.receive_tree.column("Tolerance", width=70, anchor="center")
        self.receive_tree.column("Quantity", width=65, anchor="center")
        self.receive_tree.column("QTY", width=75, anchor="center")
        self.receive_tree.column("Make", width=80, anchor="center")

        vsb_receive = ttk.Scrollbar(
            self.receive_content_container, orient="vertical", command=self.receive_tree.yview
        )
        hsb_receive = ttk.Scrollbar(
            self.receive_content_container, orient="horizontal", command=self.receive_tree.xview
        )
        self.receive_tree.configure(
            yscrollcommand=vsb_receive.set, xscrollcommand=hsb_receive.set
        )

        self.receive_tree.pack(side="top", fill="both", expand=True)
        vsb_receive.pack(side="right", fill="y")
        hsb_receive.pack(side="bottom", fill="x")

        self.receive_tree.bind("<Button-1>", self.on_receive_tree_click)

        receive_btn_bar = ttk.Frame(self.receive_content_container, padding=(0, 10, 0, 0))
        receive_btn_bar.pack(fill="x", side="bottom")

        ttk.Button(
            receive_btn_bar,
            text="☑ Select / Unmark All",
            command=self.toggle_select_all_receive,
        ).pack(side="left", padx=4)
        ttk.Button(
            receive_btn_bar,
            text="📊 Import Excel Sheet (.xlsx)",
            command=self.import_and_convert_excel,
        ).pack(side="left", padx=4)
        ttk.Button(
            receive_btn_bar,
            text="➕ Add Manual Row",
            command=self.open_add_material_dialog,
        ).pack(side="left", padx=4)
        ttk.Button(
            receive_btn_bar,
            text="❌ Delete Row",
            command=self.remove_selected_material,
        ).pack(side="left", padx=4)
        ttk.Button(
            receive_btn_bar,
            text="📦 Generate Received List (.xlsx)",
            command=self.send_receive_material_list,
        ).pack(side="left", padx=4)
        ttk.Button(
            receive_btn_bar,
            text="💾 Save to CSV",
            command=self.save_material_list_changes,
        ).pack(side="left", padx=4)

        # 5. Product Testing Frame
        self.test_frame = ttk.LabelFrame(
            scrollable_frame, text=" 5. Product Testing & Visual Inspection ", padding=12
        )
        self.test_frame.pack(fill="x", expand=True)

        test_btn_bar = ttk.Frame(self.test_frame)
        test_btn_bar.pack(fill="x")

        ttk.Button(
            test_btn_bar,
            text="🔍 Open Product Testing Window",
            command=self.open_product_testing_window,
        ).pack(side="left", padx=4)

        self.load_batches_into_dropdown()

    def open_product_testing_window(self):
        if not hasattr(self, "active_csv_path") or not self.active_csv_path:
            messagebox.showwarning("Batch Needed", "Please select a production batch from Section 1 first.", parent=self)
            return
        batch_folder = os.path.dirname(self.active_csv_path)
        ProductTestingWindow(
            self,
            batch_folder_path=batch_folder,
            user_data=self.user_data,
        )

    def toggle_material_section(self):
        if self.is_mat_frame_expanded:
            self.mat_content_container.pack_forget()
            self.btn_toggle_mat.config(text="▼ Expand Section")
            self.is_mat_frame_expanded = False
        else:
            self.mat_content_container.pack(fill="both", expand=True)
            self.btn_toggle_mat.config(text="▲ Fold Section")
            self.is_mat_frame_expanded = True

    def toggle_pcb_section(self):
        if self.is_pcb_frame_expanded:
            self.pcb_content_container.pack_forget()
            self.btn_toggle_pcb.config(text="▼ Expand Section")
            self.is_pcb_frame_expanded = False
        else:
            self.pcb_content_container.pack(fill="both", expand=True)
            self.btn_toggle_pcb.config(text="▲ Fold Section")
            self.is_pcb_frame_expanded = True

    def toggle_receive_section(self):
        if self.is_receive_frame_expanded:
            self.receive_content_container.pack_forget()
            self.btn_toggle_receive.config(text="▼ Expand Section")
            self.is_receive_frame_expanded = False
        else:
            self.receive_content_container.pack(fill="both", expand=True)
            self.btn_toggle_receive.config(text="▲ Fold Section")
            self.is_receive_frame_expanded = True

    def load_batches_into_dropdown(self):
        self.batch_map = {}
        batch_options = []

        pattern = os.path.join(self.new_prod_dir, "*", "info.csv")
        csv_files = glob.glob(pattern)

        for csv_file in csv_files:
            try:
                with open(csv_file, mode="r", encoding="utf-8") as f:
                    reader = csv.DictReader(f)
                    row = next(reader, None)
                    if row:
                        p_name = row.get("Product_Name", "Unknown")
                        b_num = row.get("Batch_Number", "Unknown")
                        display_str = f"{b_num} ({p_name})"

                        self.batch_map[display_str] = (csv_file, row)
                        batch_options.append(display_str)
            except Exception as e:
                print(f"Error reading {csv_file}: {e}")

        self.combo_batch["values"] = batch_options
        if batch_options:
            self.combo_batch.set(batch_options[0])
            self.on_batch_dropdown_selected(None)

    def on_batch_dropdown_selected(self, event):
        selected_key = self.combo_batch.get()
        if not selected_key or selected_key not in self.batch_map:
            return

        csv_path, row_data = self.batch_map[selected_key]
        self.active_csv_path = csv_path
        self.active_batch_row = row_data

        exp_date = row_data.get("Expected_Completion_Date", "N/A")
        start_date = row_data.get("Start_Date", "N/A")
        personnel = row_data.get("Assigned_Personnel", "None")

        self.lbl_batch_info.config(
            text=f"Start Date: {start_date}   |   Expected Completion: {exp_date}   |   Personnel: {personnel}",
            font=("Helvetica", 9, "normal"),
        )

        self.refresh_material_tree_from_data(row_data.get("Material_List", ""))

    def refresh_material_tree_from_data(self, raw_materials):
        self.mat_tree.delete(*self.mat_tree.get_children())
        self.pcb_tree.delete(*self.pcb_tree.get_children())
        self.receive_tree.delete(*self.receive_tree.get_children())
        self.current_materials = []

        if raw_materials:
            items = raw_materials.split(";")
            for idx, item_str in enumerate(items, 1):
                parts = item_str.split("|")
                if len(parts) >= 11:
                    po_val = True if parts[0] == "1" else False
                    m_dict = {
                        "po_check": po_val,
                        "pcb_check": po_val,
                        "receive_check": po_val,
                        "sr_no": parts[1] or str(idx),
                        "category": parts[2],
                        "part_name": parts[3],
                        "value": parts[4],
                        "package": parts[5],
                        "tolerance": parts[6],
                        "quantity": parts[7],
                        "qty": parts[8],
                        "rate": parts[9],
                        "make": parts[10],
                    }
                else:
                    m_dict = {
                        "po_check": False,
                        "pcb_check": False,
                        "receive_check": False,
                        "sr_no": str(idx),
                        "category": parts[0] if len(parts) > 0 else "",
                        "part_name": parts[1] if len(parts) > 1 else "",
                        "value": parts[2] if len(parts) > 2 else "",
                        "package": parts[3] if len(parts) > 3 else "",
                        "tolerance": "1%",
                        "quantity": parts[4] if len(parts) > 4 else "1",
                        "qty": "",
                        "rate": "",
                        "make": "",
                    }
                self.current_materials.append(m_dict)

        self.resequence_sr_no()
        self.redraw_material_tree()
        self.redraw_pcb_tree()
        self.redraw_receive_tree()

    def redraw_material_tree(self):
        self.mat_tree.delete(*self.mat_tree.get_children())

        for idx, m in enumerate(self.current_materials):
            chk_symbol = "☑" if m["po_check"] else "☐"
            row_tag = "evenrow" if idx % 2 == 0 else "oddrow"

            try:
                po_qty = float(m["qty"]) if m["qty"] else 0.0
                rate = float(m["rate"]) if m["rate"] else 0.0
                tot_price = f"{po_qty * rate:.2f}" if (po_qty * rate) > 0 else ""
            except ValueError:
                tot_price = ""

            self.mat_tree.insert(
                "",
                "end",
                values=(
                    chk_symbol,
                    m["sr_no"],
                    m["category"],
                    m["part_name"],
                    m["value"],
                    m["package"],
                    m["tolerance"],
                    m["quantity"],
                    m["qty"],
                    m["rate"],
                    tot_price,
                    m["make"],
                ),
                tags=(row_tag,),
            )

    def redraw_pcb_tree(self):
        self.pcb_tree.delete(*self.pcb_tree.get_children())

        for idx, m in enumerate(self.current_materials):
            chk_symbol = "☑" if m.get("pcb_check", False) else "☐"
            row_tag = "evenrow" if idx % 2 == 0 else "oddrow"

            self.pcb_tree.insert(
                "",
                "end",
                values=(
                    chk_symbol,
                    m["sr_no"],
                    m["category"],
                    m["part_name"],
                    m["value"],
                    m["package"],
                    m["tolerance"],
                    m["quantity"],
                    m["qty"],
                    m["make"],
                ),
                tags=(row_tag,),
            )

    def redraw_receive_tree(self):
        self.receive_tree.delete(*self.receive_tree.get_children())

        for idx, m in enumerate(self.current_materials):
            chk_symbol = "☑" if m.get("receive_check", False) else "☐"
            row_tag = "evenrow" if idx % 2 == 0 else "oddrow"

            self.receive_tree.insert(
                "",
                "end",
                values=(
                    chk_symbol,
                    m["sr_no"],
                    m["category"],
                    m["part_name"],
                    m["value"],
                    m["package"],
                    m["tolerance"],
                    m["quantity"],
                    m["qty"],
                    m["make"],
                ),
                tags=(row_tag,),
            )

    def toggle_select_all_po(self):
        if not self.current_materials:
            return
        all_selected = all(m["po_check"] for m in self.current_materials)
        new_state = not all_selected
        for m in self.current_materials:
            m["po_check"] = new_state
        self.redraw_material_tree()

    def toggle_select_all_pcb(self):
        if not self.current_materials:
            return
        all_selected = all(m.get("pcb_check", False) for m in self.current_materials)
        new_state = not all_selected
        for m in self.current_materials:
            m["pcb_check"] = new_state
        self.redraw_pcb_tree()

    def toggle_select_all_receive(self):
        if not self.current_materials:
            return
        all_selected = all(m.get("receive_check", False) for m in self.current_materials)
        new_state = not all_selected
        for m in self.current_materials:
            m["receive_check"] = new_state
        self.redraw_receive_tree()

    def on_material_tree_click(self, event):
        region = self.mat_tree.identify_region(event.x, event.y)
        if region == "cell":
            col = self.mat_tree.identify_column(event.x)
            if col == "#1":
                selected_item = self.mat_tree.identify_row(event.y)
                if selected_item:
                    vals = self.mat_tree.item(selected_item, "values")
                    target_sr = str(vals[1])
                    for m in self.current_materials:
                        if str(m["sr_no"]) == target_sr:
                            m["po_check"] = not m["po_check"]
                            break
                    self.redraw_material_tree()

    def on_pcb_tree_click(self, event):
        region = self.pcb_tree.identify_region(event.x, event.y)
        if region == "cell":
            col = self.pcb_tree.identify_column(event.x)
            if col == "#1":
                selected_item = self.pcb_tree.identify_row(event.y)
                if selected_item:
                    vals = self.pcb_tree.item(selected_item, "values")
                    target_sr = str(vals[1])
                    for m in self.current_materials:
                        if str(m["sr_no"]) == target_sr:
                            m["pcb_check"] = not m.get("pcb_check", False)
                            break
                    self.redraw_pcb_tree()

    def on_receive_tree_click(self, event):
        region = self.receive_tree.identify_region(event.x, event.y)
        if region == "cell":
            col = self.receive_tree.identify_column(event.x)
            if col == "#1":
                selected_item = self.receive_tree.identify_row(event.y)
                if selected_item:
                    vals = self.receive_tree.item(selected_item, "values")
                    target_sr = str(vals[1])
                    for m in self.current_materials:
                        if str(m["sr_no"]) == target_sr:
                            m["receive_check"] = not m.get("receive_check", False)
                            break
                    self.redraw_receive_tree()

    def import_and_convert_excel(self):
        if not HAS_OPENPYXL:
            messagebox.showerror(
                "Missing Library",
                "openpyxl is required.\nInstall with: pip install openpyxl",
                parent=self,
            )
            return

        if not hasattr(self, "active_csv_path"):
            messagebox.showwarning(
                "Warning", "Please select a production batch first.", parent=self
            )
            return

        file_path = filedialog.askopenfilename(
            title="Select Production Excel File",
            filetypes=[("Excel Files", "*.xlsx *.xls")],
            parent=self,
        )

        if not file_path:
            return

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            imported_count = 0
            current_category = ""

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]

                for row in sheet.iter_rows(values_only=True):
                    if not row or all(v is None for v in row):
                        continue

                    r_sr = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
                    r_cat = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
                    r_part = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
                    r_val = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
                    r_pkg = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
                    r_tol = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
                    r_qty = str(row[6]).strip() if len(row) > 6 and row[6] is not None else ""
                    r_po_qty = str(row[7]).strip() if len(row) > 7 and row[7] is not None else ""
                    r_rate = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
                    r_make = str(row[9]).strip() if len(row) > 9 and row[9] is not None else ""

                    if r_sr.lower() in ("sr. no.", "sr no") or r_val.lower() == "value":
                        continue

                    if r_cat:
                        current_category = r_cat

                    if r_part or r_val or r_pkg:
                        self.current_materials.append({
                            "po_check": False,
                            "pcb_check": False,
                            "receive_check": False,
                            "sr_no": "",
                            "category": current_category,
                            "part_name": r_part if r_part != "None" else "",
                            "value": r_val if r_val != "None" else "",
                            "package": r_pkg if r_pkg != "None" else "",
                            "tolerance": r_tol if r_tol != "None" else "1%",
                            "quantity": r_qty if r_qty != "None" else "",
                            "qty": r_po_qty if r_po_qty != "None" else "",
                            "rate": r_rate if r_rate != "None" else "",
                            "make": r_make if r_make != "None" else "",
                        })
                        imported_count += 1

            self.resequence_sr_no()
            self.redraw_material_tree()
            self.redraw_pcb_tree()
            self.redraw_receive_tree()
            messagebox.showinfo(
                "Import Success",
                f"Imported {imported_count} items successfully!",
                parent=self,
            )

        except Exception as e:
            messagebox.showerror(
                "Import Error", f"Failed to parse Excel file:\n{e}", parent=self
            )

    def open_add_material_dialog(self):
        if not hasattr(self, "active_csv_path"):
            messagebox.showwarning(
                "Selection Needed",
                "Please select a production batch first.",
                parent=self,
            )
            return

        dialog = tk.Toplevel(self)
        dialog.title("Add Material Row")
        dialog.geometry("380x450")
        dialog.transient(self)
        dialog.grab_set()

        fields = [
            "Category",
            "Part Name",
            "Value",
            "Package",
            "Tolerance",
            "Quantity",
            "QTY",
            "Rate",
            "Make",
        ]
        entries = {}

        for idx, field in enumerate(fields):
            ttk.Label(dialog, text=f"{field}:").grid(
                row=idx, column=0, sticky="w", padx=15, pady=4
            )
            ent = ttk.Entry(dialog, width=25)
            ent.grid(row=idx, column=1, padx=15, pady=4)
            entries[field] = ent

        def add_item():
            self.current_materials.append({
                "po_check": False,
                "pcb_check": False,
                "receive_check": False,
                "sr_no": str(len(self.current_materials) + 1),
                "category": entries["Category"].get().strip(),
                "part_name": entries["Part Name"].get().strip(),
                "value": entries["Value"].get().strip(),
                "package": entries["Package"].get().strip(),
                "tolerance": entries["Tolerance"].get().strip(),
                "quantity": entries["Quantity"].get().strip(),
                "qty": entries["QTY"].get().strip(),
                "rate": entries["Rate"].get().strip(),
                "make": entries["Make"].get().strip(),
            })
            self.resequence_sr_no()
            self.redraw_material_tree()
            self.redraw_pcb_tree()
            self.redraw_receive_tree()
            dialog.destroy()

        ttk.Button(dialog, text="Add Row", command=add_item).grid(
            row=len(fields), column=0, columnspan=2, pady=15
        )

    def remove_selected_material(self):
        selected_mat = self.mat_tree.selection()
        selected_pcb = self.pcb_tree.selection()
        selected_receive = self.receive_tree.selection()
        
        idx_to_remove = None
        if selected_mat:
            idx_to_remove = self.mat_tree.index(selected_mat[0])
        elif selected_pcb:
            idx_to_remove = self.pcb_tree.index(selected_pcb[0])
        elif selected_receive:
            idx_to_remove = self.receive_tree.index(selected_receive[0])

        if idx_to_remove is None:
            messagebox.showwarning(
                "Selection Needed",
                "Please select a row from a table to delete.",
                parent=self,
            )
            return

        del self.current_materials[idx_to_remove]
        self.resequence_sr_no()
        self.redraw_material_tree()
        self.redraw_pcb_tree()
        self.redraw_receive_tree()

    def save_material_list_changes(self):
        if not hasattr(self, "active_csv_path"):
            messagebox.showwarning(
                "Selection Needed",
                "Please select a production batch first.",
                parent=self,
            )
            return

        self.resequence_sr_no()
        mat_str_list = []
        for m in self.current_materials:
            po_flag = "1" if m["po_check"] else "0"
            mat_str_list.append(
                f"{po_flag}|{m['sr_no']}|{m['category']}|{m['part_name']}|{m['value']}|"
                f"{m['package']}|{m['tolerance']}|{m['quantity']}|{m['qty']}|{m['rate']}|{m['make']}"
            )

        joined_materials = ";".join(mat_str_list)
        self.active_batch_row["Material_List"] = joined_materials

        try:
            with open(
                self.active_csv_path, mode="w", newline="", encoding="utf-8"
            ) as f:
                writer = csv.DictWriter(
                    f, fieldnames=list(self.active_batch_row.keys())
                )
                writer.writeheader()
                writer.writerow(self.active_batch_row)

            messagebox.showinfo(
                "Data Saved",
                f"Material list saved successfully!\n\nSaved location:\n{self.active_csv_path}",
                parent=self,
            )
        except Exception as e:
            messagebox.showerror(
                "Save Error", f"Failed to save CSV:\n{e}", parent=self
            )

    def send_purchase_order(self):
        if not HAS_OPENPYXL:
            messagebox.showerror(
                "Missing Library",
                "openpyxl is required to generate Excel files.\nInstall with: pip install openpyxl",
                parent=self,
            )
            return

        if not hasattr(self, "active_csv_path"):
            messagebox.showwarning(
                "Selection Needed", "Please select a batch first.", parent=self
            )
            return

        po_items = [m for m in self.current_materials if m["po_check"]]
        if not po_items:
            messagebox.showwarning(
                "No Items Selected",
                "Please select at least one item using the [ PO ] checkbox.",
                parent=self,
            )
            return

        try:
            batch_folder = os.path.dirname(self.active_csv_path)
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"Purchase_Order_{timestamp}.xlsx"
            file_path = os.path.join(batch_folder, filename)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Purchase Order"

            current_user = getattr(auth_manager, "current_user", {})
            user_uname = "Admin"
            user_fullname = ""
            user_mobile = ""

            if isinstance(current_user, dict):
                user_uname = current_user.get("username", "Admin")
                user_fullname = current_user.get("full_name", user_uname)
                user_mobile = current_user.get("mobile_number", "")
            elif hasattr(current_user, "username"):
                user_uname = getattr(current_user, "username", "Admin")
                user_fullname = getattr(current_user, "full_name", user_uname)
                user_mobile = getattr(current_user, "mobile_number", "")

            batch_num = self.active_batch_row.get("Batch_Number", "N/A")
            product_name = self.active_batch_row.get("Product_Name", "N/A")
            assigned_personnel = self.active_batch_row.get(
                "Assigned_Personnel", "N/A"
            )
            start_date = self.active_batch_row.get("Start_Date", "N/A")
            exp_date = self.active_batch_row.get(
                "Expected_Completion_Date", "N/A"
            )

            ws.append(["PURCHASE ORDER REPORT"])

            meta_rows = [
                ("Generated Date:", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                ("Generated By User:", user_uname),
                ("Full Name:", user_fullname),
                ("Mobile Number:", user_mobile),
                ("Batch Number:", batch_num),
                ("Product Name:", product_name),
                ("Assigned Personnel:", assigned_personnel),
            ]

            for label, val in meta_rows:
                ws.append([label, val])

            ws.append(["Start Date:", start_date, "Expected Completion:", exp_date])
            ws.append([])

            headers = [
                "Sr. No.",
                "Category",
                "Part Name",
                "Value",
                "Package",
                "Tolerance",
                "Quantity",
                "PO QTY",
                "Rate",
                "Total Price",
                "Make",
            ]
            ws.append(headers)
            header_row_idx = ws.max_row
            data_start_row = header_row_idx + 1

            for idx, m in enumerate(po_items, start=1):
                raw_qty = m.get("quantity", "")
                raw_po_qty = m.get("qty", "")
                raw_rate = m.get("rate", "")

                try:
                    num_qty = float(raw_qty) if raw_qty != "" else ""
                except ValueError:
                    num_qty = raw_qty

                try:
                    num_po_qty = float(raw_po_qty) if raw_po_qty != "" else ""
                except ValueError:
                    num_po_qty = raw_po_qty

                try:
                    num_rate = float(raw_rate) if raw_rate != "" else ""
                except ValueError:
                    num_rate = raw_rate

                if isinstance(num_po_qty, (int, float)) and isinstance(num_rate, (int, float)):
                    tot_price_val = round(num_po_qty * num_rate, 2)
                else:
                    tot_price_val = ""

                ws.append([
                    str(idx),
                    m.get("category", ""),
                    m.get("part_name", ""),
                    m.get("value", ""),
                    m.get("package", ""),
                    m.get("tolerance", ""),
                    num_qty,
                    num_po_qty,
                    num_rate,
                    tot_price_val,
                    m.get("make", ""),
                ])

            data_end_row = ws.max_row

            ws.append([
                "Total Summary",
                "",
                "",
                "",
                "",
                "",
                f"=SUM(G{data_start_row}:G{data_end_row})",
                f"=SUM(H{data_start_row}:H{data_end_row})",
                "",
                f"=SUM(J{data_start_row}:J{data_end_row})",
                "",
            ])
            summary_row_idx = ws.max_row

            ws.merge_cells("A1:K1")
            ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

            bold_font = Font(bold=True)
            for r in range(2, header_row_idx - 1):
                ws[f"A{r}"].font = bold_font
                if ws[f"C{r}"].value:
                    ws[f"C{r}"].font = bold_font

            tbl_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=header_row_idx, column=col_idx)
                cell.font = bold_font
                cell.fill = tbl_header_fill
                cell.alignment = Alignment(horizontal="center")

            for row in range(data_start_row, data_end_row + 1):
                ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
                ws.cell(row=row, column=7).alignment = Alignment(horizontal="right")
                ws.cell(row=row, column=8).alignment = Alignment(horizontal="right")
                ws.cell(row=row, column=9).alignment = Alignment(horizontal="right")
                ws.cell(row=row, column=10).alignment = Alignment(horizontal="right")

            summary_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            thin_border = Border(top=Side(style="thin"), bottom=Side(style="double"))
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=summary_row_idx, column=col_idx)
                cell.font = bold_font
                cell.fill = summary_fill
                cell.border = thin_border
                if col_idx in (7, 8, 10):
                    cell.alignment = Alignment(horizontal="right")

            ws.merge_cells(f"A{summary_row_idx}:F{summary_row_idx}")
            ws[f"A{summary_row_idx}"].alignment = Alignment(horizontal="right")

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(file_path)

            messagebox.showinfo(
                "Purchase Order Generated",
                f"Purchase Order Excel sheet created successfully!\n\nSaved location:\n{file_path}",
                parent=self,
            )

        except Exception as e:
            messagebox.showerror(
                "Export Error",
                f"Failed to generate Purchase Order Excel:\n{e}",
                parent=self,
            )

    def send_pcb_material_list(self):
        if not HAS_OPENPYXL:
            messagebox.showerror("Missing Library", "openpyxl is required.", parent=self)
            return

        if not hasattr(self, "active_csv_path"):
            messagebox.showwarning("Selection Needed", "Please select a batch first.", parent=self)
            return

        pcb_items = [m for m in self.current_materials if m.get("pcb_check", False)]
        if not pcb_items:
            messagebox.showwarning("No Items Selected", "Please select at least one item using the [ Sel ] checkbox.", parent=self)
            return

        try:
            batch_folder = os.path.dirname(self.active_csv_path)
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            file_path = os.path.join(batch_folder, f"PCB_Mounting_Material_List_{timestamp}.xlsx")

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "PCB Material List"

            ws.append(["PCB MOUNTING MATERIAL LIST"])
            headers = ["Sr. No.", "Category", "Part Name", "Value", "Package", "Tolerance", "Quantity", "PO QTY", "Make"]
            ws.append([])
            ws.append(headers)
            header_row = ws.max_row
            data_start = header_row + 1

            for idx, m in enumerate(pcb_items, start=1):
                ws.append([
                    str(idx),
                    m.get("category", ""),
                    m.get("part_name", ""),
                    m.get("value", ""),
                    m.get("package", ""),
                    m.get("tolerance", ""),
                    m.get("quantity", ""),
                    m.get("qty", ""),
                    m.get("make", ""),
                ])

            data_end = ws.max_row
            ws.append([
                "Total Summary",
                "",
                "",
                "",
                "",
                "",
                f"=SUM(G{data_start}:G{data_end})",
                f"=SUM(H{data_start}:H{data_end})",
                "",
            ])
            summary_row = ws.max_row

            ws.merge_cells("A1:I1")
            ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

            bold_font = Font(bold=True)
            tbl_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=header_row, column=col_idx)
                cell.font = bold_font
                cell.fill = tbl_header_fill
                cell.alignment = Alignment(horizontal="center")

            summary_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            thin_border = Border(top=Side(style="thin"), bottom=Side(style="double"))
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=summary_row, column=col_idx)
                cell.font = bold_font
                cell.fill = summary_fill
                cell.border = thin_border
                if col_idx in (7, 8):
                    cell.alignment = Alignment(horizontal="right")

            ws.merge_cells(f"A{summary_row}:F{summary_row}")
            ws[f"A{summary_row}"].alignment = Alignment(horizontal="right")

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(file_path)
            messagebox.showinfo("Success", f"PCB Material List generated successfully:\n{file_path}", parent=self)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to generate Excel:\n{e}", parent=self)

    def send_receive_material_list(self):
        if not HAS_OPENPYXL:
            messagebox.showerror("Missing Library", "openpyxl is required.", parent=self)
            return

        if not hasattr(self, "active_csv_path"):
            messagebox.showwarning("Selection Needed", "Please select a batch first.", parent=self)
            return

        receive_items = [m for m in self.current_materials if m.get("receive_check", False)]
        if not receive_items:
            messagebox.showwarning("No Items Selected", "Please select at least one item using the [ Rcv ] checkbox.", parent=self)
            return

        try:
            batch_folder = os.path.dirname(self.active_csv_path)
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            file_path = os.path.join(batch_folder, f"Received_from_PCB_Mounting_{timestamp}.xlsx")

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Received List"

            ws.append(["RECEIVED FROM PCB MOUNTING LIST"])
            headers = ["Sr. No.", "Category", "Part Name", "Value", "Package", "Tolerance", "Quantity", "PO QTY", "Make"]
            ws.append([])
            ws.append(headers)
            header_row = ws.max_row
            data_start = header_row + 1

            for idx, m in enumerate(receive_items, start=1):
                ws.append([
                    str(idx),
                    m.get("category", ""),
                    m.get("part_name", ""),
                    m.get("value", ""),
                    m.get("package", ""),
                    m.get("tolerance", ""),
                    m.get("quantity", ""),
                    m.get("qty", ""),
                    m.get("make", ""),
                ])

            data_end = ws.max_row
            ws.append([
                "Total Summary",
                "",
                "",
                "",
                "",
                "",
                f"=SUM(G{data_start}:G{data_end})",
                f"=SUM(H{data_start}:H{data_end})",
                "",
            ])
            summary_row = ws.max_row

            ws.merge_cells("A1:I1")
            ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

            bold_font = Font(bold=True)
            tbl_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=header_row, column=col_idx)
                cell.font = bold_font
                cell.fill = tbl_header_fill
                cell.alignment = Alignment(horizontal="center")

            summary_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            thin_border = Border(top=Side(style="thin"), bottom=Side(style="double"))
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=summary_row, column=col_idx)
                cell.font = bold_font
                cell.fill = summary_fill
                cell.border = thin_border
                if col_idx in (7, 8):
                    cell.alignment = Alignment(horizontal="right")

            ws.merge_cells(f"A{summary_row}:F{summary_row}")
            ws[f"A{summary_row}"].alignment = Alignment(horizontal="right")

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(file_path)
            messagebox.showinfo("Success", f"Received List generated successfully:\n{file_path}", parent=self)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to generate Excel:\n{e}", parent=self)

    # ------------------------------------------------------------------
    # NEW PRODUCTION START
    # ------------------------------------------------------------------

    def setup_new_production_ui(self):
        self.clear_view()
        self.attached_file_path = None

        now = datetime.datetime.now()
        self.start_datetime_str = now.strftime("%Y-%m-%d %H:%M:%S")
        batch_timestamp = now.strftime("%Y%m%d_%H%M%S")
        self.batch_number = f"BATCH-{batch_timestamp}"

        top_bar = ttk.Frame(self, padding=10)
        top_bar.pack(fill="x")

        ttk.Button(
            top_bar, text="⬅ Back to Option", command=self.show_mode_selection
        ).pack(side="left")
        ttk.Label(
            top_bar,
            text="New Production Process",
            font=("Helvetica", 13, "bold"),
        ).pack(side="left", padx=15)

        main_canvas = tk.Canvas(self, borderwidth=0, highlightthickness=0)
        scrollbar = ttk.Scrollbar(
            self, orient="vertical", command=main_canvas.yview
        )
        scrollable_frame = ttk.Frame(main_canvas, padding=10)

        scrollable_frame.bind(
            "<Configure>",
            lambda e: main_canvas.configure(
                scrollregion=main_canvas.bbox("all")
            ),
        )
        main_canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        main_canvas.configure(yscrollcommand=scrollbar.set)

        main_canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        meta_frame = ttk.LabelFrame(
            scrollable_frame, text=" 1. Schedule & Details ", padding=12
        )
        meta_frame.pack(fill="x", padx=10, pady=5)

        ttk.Label(
            meta_frame,
            text=f"Batch Number: {self.batch_number}",
            font=("Helvetica", 10, "bold"),
        ).grid(row=0, column=0, sticky="w", padx=10, pady=4)
        ttk.Label(
            meta_frame,
            text=f"Start Date: {self.start_datetime_str}",
            font=("Helvetica", 10),
        ).grid(row=0, column=1, sticky="w", padx=10, pady=4)

        ttk.Label(meta_frame, text="Expected Completion Date:").grid(
            row=1, column=0, sticky="w", padx=10, pady=6
        )
        self.entry_exp_date = ttk.Entry(meta_frame, width=20)
        self.entry_exp_date.insert(
            0,
            (datetime.date.today() + datetime.timedelta(days=7)).strftime(
                "%Y-%m-%d"
            ),
        )
        self.entry_exp_date.grid(row=1, column=1, sticky="w", padx=10, pady=6)

        ttk.Label(meta_frame, text="Meeting Date:").grid(
            row=2, column=0, sticky="w", padx=10, pady=6
        )
        self.entry_meeting_date = ttk.Entry(meta_frame, width=20)
        self.entry_meeting_date.insert(
            0, datetime.date.today().strftime("%Y-%m-%d")
        )
        self.entry_meeting_date.grid(
            row=2, column=1, sticky="w", padx=10, pady=6
        )

        prod_frame = ttk.LabelFrame(
            scrollable_frame, text=" 2. Select or Add Product ", padding=12
        )
        prod_frame.pack(fill="x", padx=10, pady=8)

        ttk.Label(prod_frame, text="Select Product:").grid(
            row=0, column=0, sticky="w", padx=5, pady=5
        )
        self.load_products_from_csv()
        self.product_combo = ttk.Combobox(
            prod_frame, width=35, state="readonly", values=self.products
        )
        self.product_combo.grid(row=0, column=1, sticky="w", padx=5, pady=5)
        if self.products:
            self.product_combo.set(self.products[0])

        ttk.Button(
            prod_frame,
            text="+ Add New Product",
            command=self.open_add_product_dialog,
        ).grid(row=0, column=2, padx=10, pady=5)

        person_frame = ttk.LabelFrame(
            scrollable_frame,
            text=" 3. Assigned Personnel (Admin Tools Integration) ",
            padding=12,
        )
        person_frame.pack(fill="x", padx=10, pady=8)

        users = auth_manager.load_users()
        self.selected_personnel = {}

        if users:
            col_idx = 0
            row_idx = 0
            for u in users:
                if isinstance(u, dict):
                    uname = u.get("username", "")
                    fname = u.get("full_name") or uname
                    desig = u.get("designation", "")
                    mobile = u.get("mobile_number", "")

                    display_parts = [fname]
                    if desig:
                        display_parts.append(f"({desig})")
                    if mobile:
                        display_parts.append(f"- {mobile}")
                    display_label = " ".join(display_parts)
                else:
                    uname = getattr(u, "username", str(u))
                    fname = getattr(u, "full_name", uname) or uname
                    desig = getattr(u, "designation", "")
                    mobile = getattr(u, "mobile_number", "")

                    display_parts = [fname]
                    if desig:
                        display_parts.append(f"({desig})")
                    if mobile:
                        display_parts.append(f"- {mobile}")
                    display_label = " ".join(display_parts)

                if uname:
                    var = tk.BooleanVar(value=False)
                    chk = ttk.Checkbutton(
                        person_frame, text=display_label, variable=var
                    )
                    chk.grid(
                        row=row_idx,
                        column=col_idx,
                        sticky="w",
                        padx=12,
                        pady=4,
                    )
                    self.selected_personnel[display_label] = var

                    col_idx += 1
                    if col_idx > 1:
                        col_idx = 0
                        row_idx += 1

        doc_frame = ttk.LabelFrame(
            scrollable_frame,
            text=" 4. Production Notes & Documents ",
            padding=12,
        )
        doc_frame.pack(fill="x", padx=10, pady=8)

        ttk.Label(doc_frame, text="Notes:").pack(anchor="w", pady=(0, 4))
        self.txt_notes = tk.Text(doc_frame, height=4, font=("Helvetica", 9))
        self.txt_notes.pack(fill="x", pady=(0, 10))

        file_bar = ttk.Frame(doc_frame)
        file_bar.pack(fill="x")

        ttk.Button(
            file_bar,
            text="📎 Attach Word or Excel File",
            command=self.attach_doc_file,
        ).pack(side="left", padx=5)
        self.lbl_file_status = ttk.Label(
            file_bar, text="No file attached", font=("Helvetica", 9, "italic")
        )
        self.lbl_file_status.pack(side="left", padx=10)

        action_frame = ttk.Frame(scrollable_frame, padding=10)
        action_frame.pack(fill="x", pady=10)

        ttk.Button(
            action_frame,
            text="💾 Save Production Batch",
            command=self.start_production,
        ).pack(side="right", padx=10)

    def attach_doc_file(self):
        file_path = filedialog.askopenfilename(
            title="Select File",
            filetypes=[
                ("Documents & Spreadsheets", "*.docx *.xlsx *.doc *.xls"),
                ("All Files", "*.*"),
            ],
            parent=self,
        )
        if file_path:
            self.attached_file_path = file_path
            self.lbl_file_status.config(
                text=f"Attached: {os.path.basename(file_path)}"
            )

    def start_production(self):
        selected_product = self.product_combo.get().strip()

        if not selected_product:
            messagebox.showwarning(
                "Warning",
                "Please select or add a product before saving.",
                parent=self,
            )
            return

        assigned_users = [
            user for user, var in self.selected_personnel.items() if var.get()
        ]
        exp_date = self.entry_exp_date.get().strip()
        meeting_date = self.entry_meeting_date.get().strip()
        notes = self.txt_notes.get("1.0", tk.END).strip().replace("\n", " ")

        try:
            self.save_product_to_csv(selected_product)

            clean_prod = "".join(
                c
                for c in selected_product
                if c.isalnum() or c in (" ", "_", "-")
            ).rstrip()
            batch_folder_name = f"{self.batch_number}_{clean_prod}"
            batch_folder_path = os.path.join(
                self.new_prod_dir, batch_folder_name
            )
            os.makedirs(batch_folder_path, exist_ok=True)

            copied_file_name = "None"
            if self.attached_file_path and os.path.exists(
                self.attached_file_path
            ):
                dest_file = os.path.join(
                    batch_folder_path, os.path.basename(self.attached_file_path)
                )
                shutil.copy(self.attached_file_path, dest_file)
                copied_file_name = os.path.basename(dest_file)

            batch_csv_path = os.path.join(batch_folder_path, "info.csv")
            with open(
                batch_csv_path, mode="w", newline="", encoding="utf-8"
            ) as f:
                writer = csv.writer(f)
                writer.writerow([
                    "Batch_Number",
                    "Product_Name",
                    "Start_Date",
                    "Expected_Completion_Date",
                    "Meeting_Date",
                    "Assigned_Personnel",
                    "Attached_Document",
                    "Notes",
                    "Material_List",
                ])
                writer.writerow([
                    self.batch_number,
                    selected_product,
                    self.start_datetime_str,
                    exp_date,
                    meeting_date,
                    ";".join(assigned_users),
                    copied_file_name,
                    notes,
                    "",
                ])

            messagebox.showinfo(
                "Success",
                f"Production Batch Saved!\n\nFolder: Production/New Production Start/{batch_folder_name}",
                parent=self,
            )
            self.show_mode_selection()

        except Exception as e:
            messagebox.showerror(
                "Save Error", f"Failed to save production:\n{e}", parent=self
            )

    def open_add_product_dialog(self):
        dialog = tk.Toplevel(self)
        dialog.title("Add New Product")
        dialog.geometry("360x160")
        dialog.transient(self)
        dialog.grab_set()

        ttk.Label(dialog, text="Product Name / Model:").pack(
            anchor="w", padx=15, pady=(15, 5)
        )
        entry_name = ttk.Entry(dialog, width=35)
        entry_name.pack(padx=15, pady=5)
        entry_name.focus_set()

        def save_product():
            p_name = entry_name.get().strip()
            if p_name:
                self.save_product_to_csv(p_name)
                self.product_combo["values"] = self.products
                self.product_combo.set(p_name)
                dialog.destroy()
            else:
                messagebox.showwarning(
                    "Input Error",
                    "Please enter a product name.",
                    parent=dialog,
                )

        ttk.Button(dialog, text="Save Product", command=save_product).pack(
            pady=15
        )