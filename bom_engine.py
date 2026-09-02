import csv
from datetime import datetime
import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import config


def ensure_files_exist():
    """Ensure all default CSV files and output folders exist upon startup."""
    os.makedirs("reports", exist_ok=True)

    # Ensure booster products file exists
    if not os.path.exists(config.BOOSTER_PRODUCTS_CSV):
        try:
            with open(
                config.BOOSTER_PRODUCTS_CSV,
                mode="w",
                newline="",
                encoding="utf-8-sig",
            ) as file:
                writer = csv.writer(file)
                writer.writerow(config.PRODUCT_COLUMNS)
                writer.writerow([
                    "10.0",
                    "2",
                    "1",
                    "With Bypass",
                    "Indoor",
                    "600x400",
                    "Industrial",
                    "Yes",
                    "Yes",
                    "Yes",
                    "1250.00",
                ])
        except Exception as e:
            print(f"Could not create booster products CSV: {e}")

    # Ensure water meter products file exists
    if not os.path.exists(config.WATER_METER_PRODUCTS_CSV):
        try:
            with open(
                config.WATER_METER_PRODUCTS_CSV,
                mode="w",
                newline="",
                encoding="utf-8-sig",
            ) as file:
                writer = csv.writer(file)
                writer.writerow([
                    "device_name",
                    "tank_height_position",
                    "solenoid_valve",
                    "sensor_type",
                    "communication_medium",
                    "meter_size",
                    "price",
                ])
                writer.writerow([
                    "Water Level Sensor",
                    "Overhead Tank",
                    "Yes",
                    "Float Sensor",
                    "LoRa",
                    "DN50",
                    "4200.00",
                ])
        except Exception as e:
            print(f"Could not create water meter products CSV: {e}")

    # Ensure regular customers file exists (kept for compatibility)
    if not os.path.exists(config.CUSTOMERS_CSV):
        try:
            with open(
                config.CUSTOMERS_CSV,
                mode="w",
                newline="",
                encoding="utf-8-sig",
            ) as file:
                writer = csv.writer(file)
                writer.writerow(config.CUSTOMERS_HEADERS)
                writer.writerow(["Standard Customer", "0.0"])
        except Exception as e:
            print(f"Could not create default customers CSV: {e}")

    # Ensure product-specific customer files exist
    for path in (config.BOOSTER_CUSTOMERS_CSV, config.WATER_METER_CUSTOMERS_CSV):
        if not os.path.exists(path):
            try:
                with open(path, mode="w", newline="", encoding="utf-8-sig") as file:
                    writer = csv.writer(file)
                    writer.writerow(config.CUSTOMERS_HEADERS)
                    writer.writerow(["Standard Customer", "0.0"])
            except Exception as e:
                print(f"Could not create customer CSV {path}: {e}")

    # Ensure customers_detailed.csv exists
    crm_path = getattr(
        config,
        "CUSTOMERS_DETAILED_CSV",
        os.path.join(config.SCRIPT_DIR, "customers_detailed.csv"),
    )
    crm_headers = getattr(
        config,
        "CUSTOMERS_DETAILED_HEADERS",
        [
            "company_name",
            "website",
            "contact_number",
            "address",
            "location",
            "note",
            "call_conversion_time",
            "company_data_sent",
            "enquiry_received",
            "communication_details",
            "meeting_schedule_time",
            "meeting_agenda",
            "meeting_completed_details",
        ],
    )
    if not os.path.exists(crm_path):
        try:
            with open(
                crm_path, mode="w", newline="", encoding="utf-8-sig"
            ) as file:
                writer = csv.writer(file)
                writer.writerow(crm_headers)
        except Exception as e:
            print(f"Could not create detailed customers CRM CSV: {e}")


def read_csv_data(file_path):
    """Read any target CSV file into a list of dictionaries."""
    rows = []
    if os.path.exists(file_path):
        try:
            with open(
                file_path, mode="r", newline="", encoding="utf-8-sig"
            ) as file:
                reader = csv.DictReader(file)
                for row in reader:
                    rows.append(row)
        except Exception as e:
            print(f"Failed to read CSV file {file_path}: {e}")
    return rows


def fetch_item_dp(category, sub_cat=None, capacity=None):
    """Dynamically finds Item Name and Dealer Price (DP) from price_list_clean.csv."""
    price_list_path = getattr(config, "PRICE_LIST_CSV", config.PRODUCTS_CSV)
    if not os.path.exists(price_list_path):
        return f"{category} {sub_cat or ''}".strip(), 0.0

    try:
        with open(price_list_path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                cat_match = (
                    str(row.get("Category", "")).strip().lower()
                    == str(category).strip().lower()
                )
                sub_match = True
                if sub_cat is not None:
                    sub_match = (
                        str(row.get("SUB Category", "")).strip().lower()
                        == str(sub_cat).strip().lower()
                    )

                if cat_match and sub_match:
                    if capacity is not None and row.get("CAPACITY"):
                        try:
                            if float(row.get("CAPACITY")) != float(capacity):
                                continue
                        except ValueError:
                            if (
                                str(row.get("CAPACITY")).strip()
                                != str(capacity).strip()
                            ):
                                continue

                    csv_item_name = str(
                        row.get("Item Name", "")
                        or row.get("ITEM NAME", "")
                        or row.get("Item_Name", "")
                    ).strip()
                    if not csv_item_name:
                        csv_item_name = f"{category} {sub_cat or ''}".strip()

                    try:
                        dp_val = float(row.get("DP", 0))
                    except (ValueError, TypeError):
                        dp_val = 0.0

                    return csv_item_name, dp_val
    except Exception as e:
        print(f"Error reading item price: {e}")

    return f"{category} {sub_cat or ''}".strip(), 0.0


def fetch_next_capacity_dp(category, sub_category, target_capacity):
    """Searches price_list_clean.csv for components matching target capacity or next available size."""
    price_list_path = getattr(config, "PRICE_LIST_CSV", config.PRODUCTS_CSV)
    if not os.path.exists(price_list_path):
        return f"{category} {sub_category}".strip(), str(target_capacity), 0.0

    candidates = []
    try:
        with open(price_list_path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                cat_match = (
                    str(row.get("Category", "")).strip().lower()
                    == str(category).strip().lower()
                )
                sub_match = (
                    str(row.get("SUB Category", "")).strip().lower()
                    == str(sub_category).strip().lower()
                )

                if cat_match and sub_match:
                    cap_str = str(row.get("CAPACITY", "")).strip()
                    csv_item_name = str(
                        row.get("Item Name", "")
                        or row.get("ITEM NAME", "")
                        or row.get("Item_Name", "")
                    ).strip()
                    if not csv_item_name:
                        csv_item_name = f"{category} {sub_category}".strip()

                    try:
                        cap_val = float(cap_str)
                        dp_val = float(row.get("DP", 0))
                        candidates.append(
                            (cap_val, csv_item_name, cap_str, dp_val)
                        )
                    except ValueError:
                        continue
    except Exception as e:
        print(f"Error matching capacity: {e}")

    if not candidates:
        return f"{category} {sub_category}".strip(), str(target_capacity), 0.0

    candidates.sort(key=lambda x: x[0])

    for cap_val, item_name, cap_str, dp_val in candidates:
        if cap_val >= target_capacity:
            return item_name, cap_str, dp_val

    highest = candidates[-1]
    return highest[1], highest[2], highest[3]


def generate_bom(
    controller_type,
    num_pumps,
    num_vfd,
    pump_hp,
    olr_req,
    light_req,
    fan_qty,
    filter_qty,
    endlock_qty,
    mcb_mccb_type,
    breaker_qty,
    raw_breaker_rating,
    incomer_req,
    door_mount_req,
    total_panel_current,
):
    """Builds step-by-step Bill of Materials (BOM)."""
    bom = []

    # Step 1: Controller
    if controller_type == "AIPCU OR HMI":
        item_name, dp = fetch_item_dp("Controller", 1, capacity=1)
        bom.append({
            "Item_Name": item_name,
            "Category": "Controller",
            "SUB Category": "1",
            "Capacity": "1 Unit",
            "Qty": 1,
            "DP": dp,
        })
    else:
        item_name, dp = fetch_item_dp("Controller", 2)
        bom.append({
            "Item_Name": item_name,
            "Category": "Controller",
            "SUB Category": "2",
            "Capacity": "N/A",
            "Qty": 1,
            "DP": dp,
        })

    # Step 2: Power Supply
    item_name, ps_dp = fetch_item_dp("POWER SUPPLY", 1)
    bom.append({
        "Item_Name": item_name,
        "Category": "POWER SUPPLY",
        "SUB Category": "1",
        "Capacity": "12V/24V",
        "Qty": 1,
        "DP": ps_dp,
    })

    # Step 3: Main Switch / Load Breaker Switch
    if incomer_req == 1 and total_panel_current <= 63:
        lb_item_name, lb_cap_str, lb_dp = fetch_next_capacity_dp(
            "switch", "LOAD BREAKER", total_panel_current
        )
        bom.append({
            "Item_Name": lb_item_name,
            "Category": "switch",
            "SUB Category": "LOAD BREAKER",
            "Capacity": f"{lb_cap_str}A",
            "Qty": 1,
            "DP": lb_dp,
        })

    # Step 4: 3-Pole Door Mount Switch
    if door_mount_req == 1:
        dm_item_name, dm_cap_str, dm_dp = fetch_next_capacity_dp(
            "switch", "DOOR MOUNT", total_panel_current
        )
        if dm_item_name == "switch DOOR MOUNT":
            dm_item_name = "3 POLE SWITCH DOOR MOUNT (YELLOW) 63A"
        bom.append({
            "Item_Name": dm_item_name,
            "Category": "switch",
            "SUB Category": "DOOR MOUNT",
            "Capacity": f"{dm_cap_str}A",
            "Qty": 1,
            "DP": dm_dp,
        })

    # Step 5: Circuit Breaker
    breaker_cat = "MCCB" if "MCCB" in mcb_mccb_type.upper() else "MCB"
    breaker_sub_cat = "MCCB 3P" if breaker_cat == "MCCB" else "3P"
    b_item_name, selected_cap_str, breaker_dp = fetch_next_capacity_dp(
        breaker_cat, breaker_sub_cat, raw_breaker_rating
    )
    bom.append({
        "Item_Name": b_item_name,
        "Category": breaker_cat,
        "SUB Category": breaker_sub_cat,
        "Capacity": f"{selected_cap_str}A",
        "Qty": breaker_qty,
        "DP": breaker_dp,
    })

    # Step 6: VFD Drive Lookup
    if num_vfd > 0:
        vfd_name, vfd_cap, vfd_dp = fetch_next_capacity_dp(
            "VFD", "VFD 3 PHASE", pump_hp
        )
        bom.append({
            "Item_Name": vfd_name,
            "Category": "VFD",
            "SUB Category": "VFD 3 PHASE",
            "Capacity": f"{vfd_cap} HP",
            "Qty": num_vfd,
            "DP": vfd_dp,
        })

    # Step 7: Contactors
    contactor_qty = num_pumps * 2 if num_vfd > 0 else num_pumps
    c_name, c_cap, c_dp = fetch_next_capacity_dp(
        "CONTACTOR", "3 POLE", total_panel_current / max(1, num_pumps)
    )
    bom.append({
        "Item_Name": c_name,
        "Category": "CONTACTOR",
        "SUB Category": "3 POLE",
        "Capacity": f"{c_cap}A",
        "Qty": contactor_qty,
        "DP": c_dp,
    })

    # Step 8: Overload Relays
    if olr_req == 1:
        olr_name, olr_cap, olr_dp = fetch_next_capacity_dp(
            "OLR", "THERMAL", total_panel_current / max(1, num_pumps)
        )
        bom.append({
            "Item_Name": olr_name,
            "Category": "OLR",
            "SUB Category": "THERMAL",
            "Capacity": f"{olr_cap}A",
            "Qty": num_pumps,
            "DP": olr_dp,
        })

    # Step 9: Fans & Filters
    if fan_qty > 0:
        fan_name, fan_dp = fetch_item_dp("FAN", "COOLING")
        bom.append({
            "Item_Name": fan_name,
            "Category": "FAN",
            "SUB Category": "COOLING",
            "Capacity": "Standard",
            "Qty": fan_qty,
            "DP": fan_dp,
        })
    if filter_qty > 0:
        flt_name, flt_dp = fetch_item_dp("FILTER", "FAN FILTER")
        bom.append({
            "Item_Name": flt_name,
            "Category": "FILTER",
            "SUB Category": "FAN FILTER",
            "Capacity": "Standard",
            "Qty": filter_qty,
            "DP": flt_dp,
        })

    # Step 10: Lights & Accessories
    if light_req == 1:
        ind_name, ind_dp = fetch_item_dp("INDICATOR", "LED")
        bom.append({
            "Item_Name": ind_name,
            "Category": "INDICATOR",
            "SUB Category": "LED",
            "Capacity": "24V DC",
            "Qty": 3,
            "DP": ind_dp,
        })

    if endlock_qty > 0:
        el_name, el_dp = fetch_item_dp("ACCESSORIES", "ENDLOCK")
        bom.append({
            "Item_Name": el_name,
            "Category": "ACCESSORIES",
            "SUB Category": "ENDLOCK",
            "Capacity": "DIN Rail",
            "Qty": endlock_qty,
            "DP": el_dp,
        })

    return bom


def export_purchase_order_excel(user_data, batch_data, items_data, file_path):
    """Exports Purchase Order report including Person Name and Mobile Number."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Order"

    # Title Header
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "PURCHASE ORDER REPORT"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(
        start_color="1F497D", end_color="1F497D", fill_type="solid"
    )
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Header Metadata Block including Name and Mobile Number
    metadata = [
        ("Generated Date:", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Generated By User:", user_data.get("username", "N/A")),
        ("Full Name:", user_data.get("full_name", "N/A")),
        ("Mobile Number:", user_data.get("mobile_number", "N/A")),
        ("Batch Number:", batch_data.get("batch_number", "N/A")),
        ("Product Name:", batch_data.get("product_name", "N/A")),
        ("Assigned Personnel:", batch_data.get("assigned_personnel", "N/A")),
        ("Start Date:", batch_data.get("start_date", "N/A")),
        ("Expected Completion:", batch_data.get("expected_completion", "N/A")),
    ]

    row_idx = 2
    for label, val in metadata:
        cell_lbl = ws.cell(row=row_idx, column=1, value=label)
        cell_lbl.font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=val)
        row_idx += 1

    row_idx += 1  # Spacing row

    # Table Headers
    headers = [
        "Sr. No.",
        "Category",
        "Part Name",
        "Value",
        "Package",
        "Tolerance",
        "Quantity",
        "PO QTY",
        "Rate",
        "Make",
    ]

    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=h_text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="4F81BD", end_color="4F81BD", fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center")

    # Table Content
    row_idx += 1
    for idx, item in enumerate(items_data, 1):
        ws.cell(row=row_idx, column=1, value=idx)
        ws.cell(
            row=row_idx,
            column=2,
            value=item.get("Category", item.get("category", "")),
        )
        ws.cell(
            row=row_idx,
            column=3,
            value=item.get("Item_Name", item.get("part_name", "")),
        )
        ws.cell(
            row=row_idx, column=4, value=item.get("Capacity", item.get("value", ""))
        )
        ws.cell(row=row_idx, column=5, value=item.get("package", "N/A"))
        ws.cell(row=row_idx, column=6, value=item.get("tolerance", "N/A"))
        ws.cell(
            row=row_idx, column=7, value=item.get("Qty", item.get("quantity", 0))
        )
        ws.cell(
            row=row_idx, column=8, value=item.get("po_qty", item.get("Qty", 0))
        )
        ws.cell(
            row=row_idx, column=9, value=item.get("DP", item.get("rate", 0.0))
        )
        ws.cell(row=row_idx, column=10, value=item.get("make", "Standard"))
        row_idx += 1

    # Auto-adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(file_path)
    return file_path